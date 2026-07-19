from __future__ import annotations

import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


SCRIPT_PATH = Path(__file__).resolve().parents[1] / "scripts/02_parse_taxonomies.py"
SPEC = importlib.util.spec_from_file_location("parse_taxonomies", SCRIPT_PATH)
assert SPEC and SPEC.loader
PARSER = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = PARSER
SPEC.loader.exec_module(PARSER)


class CodeAndHierarchyTests(unittest.TestCase):
    def test_sic_numeric_code_preserves_leading_zeroes(self) -> None:
        self.assertEqual(PARSER.code_text(1110, 5), "01110")
        self.assertEqual(PARSER.code_text("01110", 5), "01110")

    def test_naics_parent_relationships_include_range_sectors(self) -> None:
        codes = {"31-33", "311", "3111", "31111", "311111"}
        self.assertEqual(PARSER.derive_naics_parent("311", codes), "31-33")
        self.assertEqual(PARSER.derive_naics_parent("3111", codes), "311")
        self.assertEqual(PARSER.derive_naics_parent("31111", codes), "3111")
        self.assertEqual(PARSER.derive_naics_parent("311111", codes), "31111")

    def test_sic_parser_preserves_exact_label_and_parents(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sic.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "reworked structure"
            worksheet.append(
                [
                    "Description",
                    "SECTION",
                    "Division",
                    "Group",
                    "Class",
                    "Sub Class",
                    "Most disaggregated level",
                    "Level headings",
                ]
            )
            worksheet.append(["Exact  Section Label  ", " A", "na", "na", "na", "na", " A", "SECTION"])
            worksheet.append(["Division", " A", "01", "na", "na", "na", "01", "Division"])
            worksheet.append(["Group", " A", "01", "011", "na", "na", "011", "Group"])
            worksheet.append(["Class", " A", "01", "011", "0111", "na", "01110", "Class"])
            worksheet.append(["Subclass", " A", "01", "011", "0111", "01111", "01111", "Sub Class"])
            workbook.create_sheet("Other sheet")
            workbook.save(path)
            manifest_input = PARSER.ManifestInput("test", path, "sic.xlsx", {})
            bundle = PARSER.parse_sic_hierarchy(manifest_input)

        section, division, group, class_record, subclass = bundle.records
        self.assertEqual(section["source_label"], "Exact  Section Label  ")
        self.assertEqual(division["source_code"], "01")
        self.assertEqual(group["parent_code"], "01")
        self.assertEqual(class_record["parent_code"], "011")
        self.assertEqual(class_record["leaf_code"], "01110")
        self.assertEqual(subclass["parent_code"], "0111")
        self.assertFalse(any(record["orphaned"] for record in bundle.records))

    def test_numerical_index_forward_fills_current_sic_code(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "activity.xlsx"
            workbook = Workbook()
            alphabetical = workbook.active
            alphabetical.title = "Alphabetical Index"
            alphabetical.append(["Title", None, None])
            alphabetical.append([None, None, None])
            alphabetical.append(["UK SIC 2007", "UK SIC 2003", "Activity"])
            alphabetical.append(["01110", "legacy", "Alphabetical activity"])
            numerical = workbook.create_sheet("Numerical Index")
            numerical.append(["Title", None, None])
            numerical.append([None, None, None])
            numerical.append(["UK SIC 2007", "UK SIC 2003", "Activity"])
            numerical.append(["01110", None, "Official class label"])
            numerical.append([None, "99999", "Granular activity"])
            numerical.append([None, "88888", "Granular activity"])
            workbook.save(path)
            context = {
                "01110": {
                    "source_code": "0111",
                    "hierarchy_path": "Section A > Division 01 > Group 01.1 > Class 01.11",
                }
            }
            manifest_input = PARSER.ManifestInput("test", path, "activity.xlsx", {})
            bundle = PARSER.parse_sic_activity_index(manifest_input, context)

        numerical_records = [
            record for record in bundle.records if record["index_type"] == "numerical"
        ]
        self.assertEqual(
            [record["sic_code"] for record in numerical_records],
            ["01110", "01110", "01110"],
        )
        self.assertEqual(numerical_records[1]["legacy_sic_2003_code"], "99999")
        self.assertTrue(all(record["hierarchy_code_matched"] for record in numerical_records))
        self.assertEqual(
            len({record["source_record_id"] for record in numerical_records}),
            len(numerical_records),
        )


class NullAndTextTests(unittest.TestCase):
    def test_suppression_markers_remain_null(self) -> None:
        self.assertIsNone(PARSER.numeric_or_none("[c]"))
        self.assertIsNone(PARSER.numeric_or_none("[z]"))
        self.assertIsNone(PARSER.integer_or_none("[c]"))
        self.assertEqual(PARSER.integer_or_none(42), 42)

    def test_naics_sections_are_extracted_without_changing_description(self) -> None:
        description = (
            "Industry definition.\n\nIllustrative Examples:\nBakery\nCafe\n\n"
            "Cross-References. Restaurants are classified elsewhere."
        )
        inclusions, exclusions = PARSER.extract_naics_sections(description)
        self.assertEqual(inclusions, "Bakery\nCafe")
        self.assertEqual(exclusions, "Restaurants are classified elsewhere.")
        self.assertTrue(description.startswith("Industry definition."))


class OSMAndRepeatabilityTests(unittest.TestCase):
    def build_aggregator(self, reverse: bool = False):
        observations = [
            ({"amenity": "restaurant", "shop": "bakery", "name": "Zulu"}, "node"),
            ({"shop": "bakery", "amenity": "restaurant", "name": "Alpha"}, "way"),
        ]
        if reverse:
            observations.reverse()
        aggregator = PARSER.OSMTagAggregator(max_examples=2)
        for tags, geometry_type in observations:
            self.assertTrue(aggregator.add(tags, geometry_type))
        return aggregator

    def test_osm_aggregation_and_filtering(self) -> None:
        aggregator = self.build_aggregator()
        self.assertFalse(aggregator.add({"amenity": "bench", "name": "Not a venue"}, "node"))
        records = aggregator.to_records("sample.osm.pbf")
        self.assertEqual(len(records), 1)
        record = records[0]
        self.assertEqual(record["primary_key"], "amenity")
        self.assertEqual(record["primary_value"], "restaurant")
        self.assertEqual(record["observed_count"], 2)
        self.assertEqual(record["example_names"], ["Alpha", "Zulu"])
        self.assertEqual(record["geometry_types"], ["node", "way"])
        self.assertEqual(record["secondary_tags_json"], '{"shop":"bakery"}')

    def test_aggregation_is_input_order_independent(self) -> None:
        forward = self.build_aggregator().to_records("sample.osm.pbf")
        reverse = self.build_aggregator(reverse=True).to_records("sample.osm.pbf")
        self.assertEqual(forward, reverse)

    def test_parquet_bytes_are_repeatable(self) -> None:
        records = self.build_aggregator().to_records("sample.osm.pbf")
        schema = PARSER.common_schema(
            [
                PARSER.pa.field("primary_key", PARSER.pa.string()),
                PARSER.pa.field("primary_value", PARSER.pa.string()),
                PARSER.pa.field("secondary_tags_json", PARSER.pa.string()),
                PARSER.pa.field("example_names", PARSER.pa.list_(PARSER.pa.string())),
                PARSER.pa.field("geometry_types", PARSER.pa.list_(PARSER.pa.string())),
            ]
        )
        stats = {
            "duplicates": 0,
            "malformed": 0,
            "unmatched": 0,
            "inspected": ["test"],
            "counts_by_level": {"amenity": 1},
            "warnings": [],
        }
        with tempfile.TemporaryDirectory() as directory:
            first = Path(directory) / "first.parquet"
            second = Path(directory) / "second.parquet"
            PARSER.write_parquet_atomic(records, schema, first, stats, "sample.osm.pbf")
            PARSER.write_parquet_atomic(records, schema, second, stats, "sample.osm.pbf")
            self.assertEqual(PARSER.sha256_file(first), PARSER.sha256_file(second))


if __name__ == "__main__":
    unittest.main()
