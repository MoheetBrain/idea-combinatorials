#!/usr/bin/env python3
"""Parse occupation ontologies, crosswalks, the BLS matrix, O*NET tasks, and exposure labels."""

from __future__ import annotations

import argparse
import csv
import html
import io
import json
import math
import os
import re
import sys
import zipfile
from collections import defaultdict
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from labour_automation_common import (
    CAPABILITY_EVIDENCE_SCHEMA,
    CROSSWALK_SCHEMA,
    INDUSTRY_OBS_SCHEMA,
    ONTOLOGY_SCHEMA,
    ROOT,
    TASK_SCHEMA,
    read_manifest,
    stable_id,
    stable_json,
    write_parquet,
)


OUTPUTS = {
    "ontology": ROOT / "data/interim/occupation_ontology_v0.parquet",
    "crosswalks": ROOT / "data/interim/occupation_crosswalks_v0.parquet",
    "industry": ROOT / "data/interim/industry_occupation_observations_v0.parquet",
    "tasks": ROOT / "data/interim/occupation_tasks_v0.parquet",
    "capability": ROOT / "data/interim/task_capability_evidence_v0.parquet",
}
PARSER_VERSION = "1"


def clean(value: Any) -> str | None:
    if value is None:
        return None
    text = " ".join(str(value).split()).strip()
    return None if not text or text.casefold() in {"<blank>", "nan", "none"} else text


def code_text(value: Any) -> str | None:
    text = clean(value)
    if text is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return text


def number(value: Any, multiplier: float = 1.0) -> float | None:
    text = clean(value)
    if text is None or text in {"—", "-", "*"}:
        return None
    try:
        result = float(text.replace(",", "")) * multiplier
        return result if math.isfinite(result) else None
    except ValueError:
        return None


def member(archive: zipfile.ZipFile, suffix: str) -> str:
    matches = [name for name in archive.namelist() if Path(name).name == suffix]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one ZIP member ending {suffix!r}, found {matches}")
    return matches[0]


def zip_dicts(archive: zipfile.ZipFile, suffix: str, delimiter: str = ",") -> Iterable[dict[str, str]]:
    handle = io.TextIOWrapper(archive.open(member(archive, suffix)), encoding="utf-8-sig", newline="")
    try:
        yield from csv.DictReader(handle, delimiter=delimiter)
    finally:
        handle.close()


def ontology_row(taxonomy: str, code: str, title: str, description: str | None, parent: str | None,
                 level: str, release: str, source_id: str, **extra: Any) -> dict[str, Any]:
    return {
        "occupation_record_id": stable_id("occ", taxonomy, code), "source_taxonomy": taxonomy,
        "source_occupation_code": code, "source_occupation_title": title,
        "source_description": description, "parent_code": parent, "hierarchy_level": level,
        "release_version": release, "source_id": source_id,
        "provenance_json": stable_json({"parser_version": PARSER_VERSION, "source_id": source_id, **extra}),
    }


def parse_ontology(paths: dict[str, Path]) -> list[dict[str, Any]]:
    rows: dict[tuple[str, str], dict[str, Any]] = {}
    with zipfile.ZipFile(paths["onet_database"]) as archive:
        for raw in zip_dicts(archive, "occupation_data.csv"):
            code, title = raw["O*NET-SOC Code"], raw["Title"]
            rows[("ONET_SOC_2019", code)] = ontology_row(
                "ONET_SOC_2019", code, title, clean(raw.get("Description")), code.split(".")[0],
                "detailed_occupation", "30.3", "onet_database", source_member="occupation_data.csv",
            )

    wb = load_workbook(paths["onet_soc_crosswalk"], read_only=True, data_only=True)
    ws = wb.worksheets[0]
    for row_number, raw in enumerate(ws.iter_rows(min_row=5, values_only=True), 5):
        soc, title = code_text(raw[2]), clean(raw[3])
        if soc and title:
            rows.setdefault(("US_SOC_2018", soc), ontology_row(
                "US_SOC_2018", soc, title, None, soc[:2] + "-0000", "detailed_occupation",
                "2018", "onet_soc_crosswalk", source_sheet=ws.title, source_row=row_number,
            ))
    wb.close()

    wb = load_workbook(paths["onet_esco_crosswalk"], read_only=True, data_only=True)
    ws = wb.worksheets[0]
    for row_number, raw in enumerate(ws.iter_rows(min_row=5, values_only=True), 5):
        esco, title = code_text(raw[0]), clean(raw[1])
        if esco and title:
            rows.setdefault(("ESCO_ISCO", esco), ontology_row(
                "ESCO_ISCO", esco, title, None, esco.split(".")[0], "occupation",
                "current", "onet_esco_crosswalk", source_sheet=ws.title, source_row=row_number,
            ))
    wb.close()

    wb = load_workbook(paths["ons_soc2020_volume1"], read_only=True, data_only=True)
    ws = wb["SOC2020 descriptions"]
    parents: dict[int, str] = {}
    levels = ["major_group", "sub_major_group", "minor_group", "unit_group"]
    for row_number, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        selected = next(((idx, code_text(raw[idx])) for idx in range(4) if clean(raw[idx])), None)
        title = clean(raw[4])
        if not selected or not title:
            continue
        idx, code = selected
        assert code is not None
        parent = parents.get(idx - 1) if idx else None
        parents[idx] = code
        for deeper in range(idx + 1, 4):
            parents.pop(deeper, None)
        description = clean(raw[6])
        rows[("UK_SOC_2020", code)] = ontology_row(
            "UK_SOC_2020", code, title, description, parent, levels[idx], "2025-12-03",
            "ons_soc2020_volume1", source_sheet=ws.title, source_row=row_number,
        )
    wb.close()

    wb = load_workbook(paths["ons_extended_soc2020"], read_only=True, data_only=True)
    ws = wb["Extended SOC framework"]
    parents = {}
    levels_ext = ["major_group", "sub_major_group", "minor_group", "unit_group", "sub_unit_group"]
    for row_number, raw in enumerate(ws.iter_rows(min_row=3, values_only=True), 3):
        selected = next(((idx, code_text(raw[idx + 1])) for idx in range(5) if clean(raw[idx + 1])), None)
        title = clean(raw[6])
        if not selected or not title:
            continue
        idx, code = selected
        assert code is not None
        parent = parents.get(idx - 1) if idx else None
        parents[idx] = code
        for deeper in range(idx + 1, 5):
            parents.pop(deeper, None)
        rows.setdefault(("UK_EXTENDED_SOC_2020", code), ontology_row(
            "UK_EXTENDED_SOC_2020", code, title, None, parent, levels_ext[idx], "V6",
            "ons_extended_soc2020", source_sheet=ws.title, source_row=row_number,
        ))
    wb.close()

    wb = load_workbook(paths["ons_soc2020_coding_index"], read_only=True, data_only=True)
    ws = wb["SOC2020 coding index"]
    for row_number, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        isco, title = code_text(raw[12]), clean(raw[20])
        if isco and title and re.fullmatch(r"\d{4}", isco):
            rows.setdefault(("ISCO_08", isco), ontology_row(
                "ISCO_08", isco, title, None, isco[:3] if len(isco) == 4 else None,
                "unit_group", "2008", "ons_soc2020_coding_index",
                source_sheet=ws.title, source_row=row_number,
            ))
    wb.close()
    return list(rows.values())


def parse_crosswalks(paths: dict[str, Path]) -> list[dict[str, Any]]:
    raw_rows: list[dict[str, Any]] = []
    wb = load_workbook(paths["onet_soc_crosswalk"], read_only=True, data_only=True)
    ws = wb.worksheets[0]
    for row_number, raw in enumerate(ws.iter_rows(min_row=5, values_only=True), 5):
        left, left_title, right, right_title = code_text(raw[0]), clean(raw[1]), code_text(raw[2]), clean(raw[3])
        if not all((left, left_title, right, right_title)):
            continue
        raw_rows.append({
            "left_taxonomy": "ONET_SOC_2019", "left_code": left, "left_title": left_title,
            "right_taxonomy": "US_SOC_2018", "right_code": right, "right_title": right_title,
            "relationship": "equivalent" if left.endswith(".00") else "narrower",
            "mapping_method": "official_crosswalk", "mapping_confidence": 0.98 if left.endswith(".00") else 0.92,
            "human_validated": True, "source_id": "onet_soc_crosswalk", "requires_manual_review": False,
            "provenance_json": stable_json({"source_sheet": ws.title, "source_row": row_number}),
        })
    wb.close()

    wb = load_workbook(paths["onet_esco_crosswalk"], read_only=True, data_only=True)
    ws = wb.worksheets[0]
    for row_number, raw in enumerate(ws.iter_rows(min_row=5, values_only=True), 5):
        left, left_title, right, right_title = code_text(raw[0]), clean(raw[1]), code_text(raw[2]), clean(raw[3])
        if not all((left, left_title, right, right_title)):
            continue
        raw_rows.append({
            "left_taxonomy": "ESCO_ISCO", "left_code": left, "left_title": left_title,
            "right_taxonomy": "ONET_SOC_2019", "right_code": right, "right_title": right_title,
            "relationship": "partial", "mapping_method": "official_ai_plus_human_validated_crosswalk",
            "mapping_confidence": 0.82, "human_validated": True, "source_id": "onet_esco_crosswalk",
            "requires_manual_review": True,
            "provenance_json": stable_json({"source_sheet": ws.title, "source_row": row_number}),
        })
    wb.close()

    pairs: dict[tuple[str, str, str, str], int] = {}
    wb = load_workbook(paths["ons_soc2020_coding_index"], read_only=True, data_only=True)
    ws = wb["SOC2020 coding index"]
    for row_number, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        left, left_title = code_text(raw[6]), clean(raw[19])
        right, right_title = code_text(raw[12]), clean(raw[20])
        if all((left, left_title, right, right_title)) and re.fullmatch(r"\d{4}", left) and re.fullmatch(r"\d{4}", right):
            pairs.setdefault((left, left_title, right, right_title), row_number)
    wb.close()
    for (left, left_title, right, right_title), row_number in pairs.items():
        raw_rows.append({
            "left_taxonomy": "UK_SOC_2020", "left_code": left, "left_title": left_title,
            "right_taxonomy": "ISCO_08", "right_code": right, "right_title": right_title,
            "relationship": "partial", "mapping_method": "official_coding_index_bridge",
            "mapping_confidence": 0.72, "human_validated": False, "source_id": "ons_soc2020_coding_index",
            "requires_manual_review": True,
            "provenance_json": stable_json({"source_sheet": ws.title, "first_source_row": row_number}),
        })
    counts: dict[tuple[str, str, str], int] = defaultdict(int)
    for row in raw_rows:
        counts[(row["left_taxonomy"], row["left_code"], row["right_taxonomy"])] += 1
    for row in raw_rows:
        row["mapping_weight"] = 1.0 / counts[(row["left_taxonomy"], row["left_code"], row["right_taxonomy"])]
    return raw_rows


class MatrixTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.in_table = False
        self.in_cell = False
        self.current_cell: list[str] = []
        self.current_row: list[str] = []
        self.rows: list[list[str]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attributes = {key.casefold(): value for key, value in attrs}
        if tag.casefold() == "table" and (attributes.get("id") or "").casefold() == "mytable":
            self.in_table = True
        elif self.in_table and tag.casefold() == "tr":
            self.current_row = []
        elif self.in_table and tag.casefold() in {"td", "th"}:
            self.in_cell = True
            self.current_cell = []

    def handle_data(self, data: str) -> None:
        if self.in_cell:
            self.current_cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if self.in_table and tag.casefold() in {"td", "th"} and self.in_cell:
            self.current_row.append(" ".join("".join(self.current_cell).split()))
            self.in_cell = False
        elif self.in_table and tag.casefold() == "tr" and self.current_row:
            self.rows.append(self.current_row)
        elif self.in_table and tag.casefold() == "table":
            self.in_table = False


def industry_level(code: str) -> str:
    digits = re.sub(r"\D", "", code)
    if re.fullmatch(r"\d{6}", code):
        return "naics_6"
    if re.fullmatch(r"\d{3,5}0+", code):
        return f"naics_{len(digits.rstrip('0'))}"
    return "matrix_composite"


def parse_bls_matrix(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with zipfile.ZipFile(path) as archive:
        index = json.loads(archive.read("matrix_index.json"))
        for item in index:
            parser = MatrixTableParser()
            parser.feed(archive.read(item["member"]).decode("utf-8", errors="replace"))
            data_rows = [row for row in parser.rows if len(row) >= 11 and re.fullmatch(r"[0-9A-Z-]{7,8}", row[1] or "")]
            total = next((number(row[3], 1000.0) for row in data_rows if row[1] == "00-0000"), None)
            for source_row, raw in enumerate(data_rows, 1):
                if raw[2].casefold() != "line item" or raw[1] == "00-0000":
                    continue
                employment = number(raw[3], 1000.0)
                share = employment / total if employment is not None and total and total > 0 else None
                if employment is None or share is None:
                    continue
                rows.append({
                    "observation_id": stable_id("nem", item["industry_code"], raw[1]),
                    "industry_code": item["industry_code"], "industry_code_level": industry_level(item["industry_code"]),
                    "industry_title": item["industry_title"], "occupation_code": raw[1],
                    "occupation_title": raw[0], "occupation_type": raw[2],
                    "occupation_employment": employment, "industry_total_employment": total,
                    "occupation_share": share, "projected_employment": number(raw[6], 1000.0),
                    "source_year": 2024, "geography": "United States", "evidence_grade": "B",
                    "confidence": 0.96, "source_id": "bls_nem_2024_2034", "source_member": item["member"],
                    "provenance_json": stable_json({
                        "parser_version": PARSER_VERSION, "source_url": item["source_url"],
                        "source_row_within_table": source_row, "published_percent_of_industry": number(raw[4]),
                        "employment_unit_original": "thousands", "share_method": "occupation employment / total industry employment",
                    }),
                })
    return rows


def parse_tasks(path: Path) -> list[dict[str, Any]]:
    with zipfile.ZipFile(path) as archive:
        ratings: dict[tuple[str, str], dict[str, Any]] = defaultdict(lambda: {"frequency": {}})
        for raw in zip_dicts(archive, "task_ratings.csv"):
            key = (raw["O*NET-SOC Code"], raw["Task ID"])
            scale, value, category = raw["Scale ID"], number(raw["Data Value"]), raw["Category"]
            if scale == "IM":
                ratings[key]["importance"] = value
            elif scale == "RT":
                ratings[key]["relevance"] = value
            elif scale == "FT" and category and value is not None:
                ratings[key]["frequency"][category] = value
        dwas: dict[tuple[str, str], list[str]] = defaultdict(list)
        for raw in zip_dicts(archive, "tasks_to_dwas.csv"):
            dwas[(raw["O*NET-SOC Code"], raw["Task ID"])].append(raw["DWA Element ID"])
        contexts: dict[str, list[str]] = defaultdict(list)
        for raw in zip_dicts(archive, "work_context.csv"):
            if raw["Scale ID"] == "CX" and not raw["Category"] and (number(raw["Data Value"]) or 0) >= 4:
                contexts[raw["O*NET-SOC Code"]].append(raw["Element ID"])
        rows = []
        for raw in zip_dicts(archive, "task_statements.csv"):
            key = (raw["O*NET-SOC Code"], raw["Task ID"])
            rating = ratings[key]
            importance, relevance = rating.get("importance"), rating.get("relevance")
            weight = (importance / 5.0) * (relevance / 100.0) if importance is not None and relevance is not None else None
            rows.append({
                "occupation_code": raw["O*NET-SOC Code"], "occupation_title": raw["Title"],
                "task_id": raw["Task ID"], "task_statement": raw["Task"], "task_type": raw["Task Type"],
                "task_importance": importance, "task_relevance": relevance,
                "task_frequency_if_available": stable_json(rating["frequency"]) if rating["frequency"] else None,
                "task_weight_proxy": weight,
                "task_weight_proxy_method": "(O*NET importance / 5) * (O*NET relevance / 100); proxy, not time share",
                "detailed_work_activity_ids": stable_json(sorted(set(dwas[key]))),
                "work_context_ids": stable_json(sorted(set(contexts[raw["O*NET-SOC Code"]]))),
                "source_release": "30.3", "source_id": "onet_database",
                "provenance_json": stable_json({
                    "parser_version": PARSER_VERSION, "task_date": raw["Date"],
                    "domain_source": raw["Domain Source"],
                    "work_context_scope": "occupation-level context; not a task-time observation",
                    "frequency_semantics": "category distribution; not converted to hours",
                }),
            })
    return rows


def task_id(value: Any) -> str | None:
    text = clean(value)
    if text and re.fullmatch(r"\d+(?:\.0)?", text):
        return text.split(".")[0]
    return text


def parse_capability(path: Path, release: str) -> list[dict[str, Any]]:
    rows: dict[tuple[str, str], dict[str, Any]] = {}
    with zipfile.ZipFile(path) as archive:
        for raw in zip_dicts(archive, "full_labelset.tsv", delimiter="\t"):
            occupation, task = clean(raw.get("O*NET-SOC Code")), task_id(raw.get("Task ID"))
            statement = clean(raw.get("Task"))
            if not occupation or not task or not statement:
                continue
            row = {
                "evidence_id": stable_id("exposure", occupation, task), "occupation_code": occupation,
                "occupation_title": clean(raw.get("Title")), "task_id": task, "task_statement": statement,
                "human_exposure_label": clean(raw.get("human_exposure_agg")),
                "model_exposure_label": clean(raw.get("gpt4_exposure")),
                "model_exposure_alt_label": clean(raw.get("gpt4_exposure_alt_rubric")),
                "published_alpha": number(raw.get("alpha")), "published_beta": number(raw.get("beta")),
                "published_gamma": number(raw.get("gamma")),
                "published_automation_label": clean(raw.get("gpt4_automation")),
                "technology_modality": "llm_exposure", "methodology_version": "GPTs are GPTs repository",
                "evidence_grade": "D", "source_id": "openai_gpts_are_gpts",
                "provenance_json": stable_json({
                    "repository_commit": release,
                    "warning": "published exposure is not task completion, adoption, headcount reduction, or payroll savings",
                }),
            }
            rows[(occupation, task)] = row
    return list(rows.values())


def run(force: bool, skip_existing: bool) -> int:
    if skip_existing and all(path.exists() for path in OUTPUTS.values()) and not force:
        print("All occupation-source outputs already exist; skipped")
        return 0
    manifest = read_manifest()
    required = {
        "onet_database", "onet_soc_crosswalk", "onet_esco_crosswalk", "bls_nem_2024_2034",
        "ons_soc2020_volume1", "ons_extended_soc2020", "ons_soc2020_coding_index", "openai_gpts_are_gpts",
    }
    missing = required - manifest.keys()
    if missing:
        raise RuntimeError(f"Occupation manifest missing {sorted(missing)}")
    paths = {source_id: ROOT / manifest[source_id]["local_path"] for source_id in required}
    ontology = parse_ontology(paths)
    crosswalks = parse_crosswalks(paths)
    industry = parse_bls_matrix(paths["bls_nem_2024_2034"])
    tasks = parse_tasks(paths["onet_database"])
    capability = parse_capability(paths["openai_gpts_are_gpts"], manifest["openai_gpts_are_gpts"]["release_version"])
    if not all((ontology, crosswalks, industry, tasks, capability)):
        raise RuntimeError("One or more required parsed tables are empty")
    counts = {
        "ontology": write_parquet(OUTPUTS["ontology"], ontology, ONTOLOGY_SCHEMA, ["source_taxonomy", "source_occupation_code"]),
        "crosswalks": write_parquet(OUTPUTS["crosswalks"], crosswalks, CROSSWALK_SCHEMA, ["left_taxonomy", "left_code", "right_taxonomy", "right_code"]),
        "industry": write_parquet(OUTPUTS["industry"], industry, INDUSTRY_OBS_SCHEMA, ["industry_code", "occupation_code"]),
        "tasks": write_parquet(OUTPUTS["tasks"], tasks, TASK_SCHEMA, ["occupation_code", "task_id"]),
        "capability": write_parquet(OUTPUTS["capability"], capability, CAPABILITY_EVIDENCE_SCHEMA, ["occupation_code", "task_id"]),
    }
    print("Parsed " + ", ".join(f"{name}={count:,}" for name, count in counts.items()))
    return 0


def main() -> int:
    parser = argparse.ArgumentParser()
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--force", action="store_true")
    group.add_argument("--skip-existing", action="store_true")
    args = parser.parse_args()
    return run(args.force, args.skip_existing)


if __name__ == "__main__":
    raise SystemExit(main())
