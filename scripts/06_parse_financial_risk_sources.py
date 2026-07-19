#!/usr/bin/env python3
"""Parse official financial, employment, pay, demography, and insolvency data."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sys
from collections import defaultdict
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable
from zipfile import ZipFile

try:
    import pyarrow as pa
    import pyarrow.parquet as pq
    from openpyxl import load_workbook
except ModuleNotFoundError as exc:  # pragma: no cover
    raise SystemExit(f"Missing dependency {exc.name!r}; install requirements.txt") from exc


ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "data/enrichment_source_manifest.csv"
OUTPUTS = {
    "financial": ROOT / "data/interim/industry_financial_observations.parquet",
    "employment": ROOT / "data/interim/industry_employment_observations.parquet",
    "risk": ROOT / "data/interim/industry_risk_observations.parquet",
}
PARSER_VERSION = "1"
CHUNK = 1024 * 1024


FINANCIAL_SCHEMA = pa.schema([
    pa.field("observation_id", pa.string()), pa.field("source_id", pa.string()),
    pa.field("source_file", pa.string()), pa.field("source_sheet", pa.string()),
    pa.field("source_row", pa.int64()), pa.field("industry_code", pa.string()),
    pa.field("industry_code_level", pa.string()), pa.field("industry_label", pa.string()),
    pa.field("year", pa.int64()), pa.field("geography", pa.string()),
    pa.field("population_basis", pa.string()), pa.field("reporting_unit_count", pa.float64()),
    pa.field("enterprise_count", pa.float64()), pa.field("local_unit_count", pa.float64()),
    pa.field("turnover_total_gbp", pa.float64()), pa.field("purchases_total_gbp", pa.float64()),
    pa.field("employment_cost_total_gbp", pa.float64()), pa.field("agva_total_gbp", pa.float64()),
    pa.field("capital_expenditure_total_gbp", pa.float64()),
    pa.field("employment_size_band", pa.string()), pa.field("turnover_size_band", pa.string()),
    pa.field("band_count", pa.float64()), pa.field("turnover_cv", pa.float64()),
    pa.field("agva_cv", pa.float64()), pa.field("purchases_cv", pa.float64()),
    pa.field("employment_cost_cv", pa.float64()), pa.field("unit_notes", pa.string()),
    pa.field("quality_notes", pa.string()), pa.field("provenance_json", pa.string()),
])

EMPLOYMENT_SCHEMA = pa.schema([
    pa.field("observation_id", pa.string()), pa.field("source_id", pa.string()),
    pa.field("source_file", pa.string()), pa.field("source_sheet", pa.string()),
    pa.field("source_row", pa.int64()), pa.field("industry_code", pa.string()),
    pa.field("industry_code_level", pa.string()), pa.field("industry_label", pa.string()),
    pa.field("occupation_code", pa.string()), pa.field("occupation_label", pa.string()),
    pa.field("year", pa.int64()), pa.field("geography", pa.string()),
    pa.field("employment_total", pa.float64()), pa.field("employee_total", pa.float64()),
    pa.field("full_time_total", pa.float64()), pa.field("part_time_total", pa.float64()),
    pa.field("median_hourly_pay_gbp", pa.float64()), pa.field("median_annual_pay_gbp", pa.float64()),
    pa.field("median_weekly_hours", pa.float64()), pa.field("employment_cv", pa.float64()),
    pa.field("hourly_pay_cv", pa.float64()), pa.field("annual_pay_cv", pa.float64()),
    pa.field("weekly_hours_cv", pa.float64()), pa.field("population_basis", pa.string()),
    pa.field("unit_notes", pa.string()), pa.field("provenance_json", pa.string()),
])

RISK_SCHEMA = pa.schema([
    pa.field("observation_id", pa.string()), pa.field("source_id", pa.string()),
    pa.field("source_file", pa.string()), pa.field("source_sheet", pa.string()),
    pa.field("source_row", pa.int64()), pa.field("industry_code", pa.string()),
    pa.field("industry_code_level", pa.string()), pa.field("industry_label", pa.string()),
    pa.field("year", pa.int64()), pa.field("geography", pa.string()),
    pa.field("birth_count", pa.float64()), pa.field("death_count", pa.float64()),
    pa.field("active_business_stock", pa.float64()), pa.field("birth_rate", pa.float64()),
    pa.field("death_rate", pa.float64()), pa.field("survival_1y", pa.float64()),
    pa.field("survival_2y", pa.float64()), pa.field("survival_3y", pa.float64()),
    pa.field("survival_4y", pa.float64()), pa.field("survival_5y", pa.float64()),
    pa.field("survival_cohort_year", pa.int64()), pa.field("insolvency_count", pa.float64()),
    pa.field("insolvency_rate_per_10000", pa.float64()), pa.field("insolvency_type", pa.string()),
    pa.field("population_basis", pa.string()), pa.field("unit_notes", pa.string()),
    pa.field("provenance_json", pa.string()),
])


def stable_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def stable_id(prefix: str, *parts: Any) -> str:
    return f"{prefix}_{hashlib.sha256(stable_json(parts).encode()).hexdigest()[:24]}"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(CHUNK):
            digest.update(chunk)
    return digest.hexdigest()


def number(value: Any, multiplier: float = 1.0) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value) * multiplier
    text = str(value).strip().replace(",", "")
    if not text or text.casefold() in {"*", "-", "x", "[x]", "[c]", "[low]", "..", ":", ".", "n/a"}:
        return None
    try:
        return float(text) * multiplier
    except ValueError:
        return None


def code(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip().replace("(part)", "").replace("(Part)", "").strip()
    text = re.sub(r"\s+", "", text).replace(".", "")
    if re.fullmatch(r"\d{1,5}", text):
        if len(text) == 1:
            return text.zfill(2)
        return text
    if re.fullmatch(r"[A-U]", text, re.I):
        return text.upper()
    return None


def level(industry_code: str | None) -> str | None:
    if not industry_code:
        return None
    if industry_code.isalpha():
        return "section"
    return {2: "division", 3: "group", 4: "class", 5: "subclass"}.get(len(industry_code))


def split_code_label(value: Any) -> tuple[str | None, str | None]:
    if value is None:
        return None, None
    match = re.match(r"^\s*([A-U]|\d{2,5})\s*:\s*(.+?)\s*$", str(value), re.I)
    if not match:
        return None, None
    return code(match.group(1)), match.group(2).strip()


def source_file(manifest: dict[str, dict[str, str]], source_id: str) -> Path:
    row = manifest[source_id]
    path = ROOT / row["local_path"]
    if row.get("validation_status") != "passed" or not path.exists():
        raise RuntimeError(f"Manifest source {source_id} is unavailable or invalid")
    actual = sha256(path)
    if actual != row.get("sha256"):
        raise RuntimeError(f"SHA-256 mismatch for {source_id}: manifest={row.get('sha256')} local={actual}")
    return path


def load_manifest() -> dict[str, dict[str, str]]:
    with MANIFEST.open(newline="", encoding="utf-8") as handle:
        rows = {row["source_id"]: row for row in csv.DictReader(handle)}
    required = {
        "abs_2024_sections", "abs_2024_employment_size_band", "abs_2024_quality",
        "abs_2024_regional", "uk_business_2025", "bres_2024_provisional",
        "bres_2024_cv", "ashe_2025_table29", "ashe_2025_table5",
        "business_demography_2024", "business_insolvency_demography_2015_2025",
        "company_insolvencies_industry_2026_06",
    }
    missing = sorted(required - rows.keys())
    if missing:
        raise RuntimeError(f"Enrichment manifest is missing: {', '.join(missing)}")
    return rows


def provenance(source_id: str, sheet: str, row: int, **extra: Any) -> str:
    return stable_json({"parser_version": PARSER_VERSION, "source_id": source_id, "sheet": sheet, "row": row, **extra})


def financial_row(**updates: Any) -> dict[str, Any]:
    row = {field.name: None for field in FINANCIAL_SCHEMA}
    row.update(updates)
    row["observation_id"] = stable_id("finobs", row["source_id"], row["source_sheet"], row["source_row"], row["industry_code"], row["year"], row["geography"], row["employment_size_band"], row["turnover_size_band"])
    return row


def employment_row(**updates: Any) -> dict[str, Any]:
    row = {field.name: None for field in EMPLOYMENT_SCHEMA}
    row.update(updates)
    row["observation_id"] = stable_id("labobs", row["source_id"], row["source_sheet"], row["source_row"], row["industry_code"], row["occupation_code"], row["year"], row["geography"])
    return row


def risk_row(**updates: Any) -> dict[str, Any]:
    row = {field.name: None for field in RISK_SCHEMA}
    row.update(updates)
    row["observation_id"] = stable_id("riskobs", row["source_id"], row["source_sheet"], row["source_row"], row["industry_code"], row["year"], row["geography"], row["insolvency_type"])
    return row


def parse_abs(manifest: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    path = source_file(manifest, "abs_2024_sections")
    quality_path = source_file(manifest, "abs_2024_quality")
    quality: dict[str, tuple[float | None, float | None, float | None, float | None]] = {}
    qwb = load_workbook(quality_path, read_only=True, data_only=True)
    for row in qwb["Quality Measures"].iter_rows(min_row=9, values_only=True):
        sic, year = code(row[0]), number(row[2])
        if sic and year == 2024:
            quality[sic] = tuple(number(row[i]) for i in range(7, 11))  # type: ignore[assignment]
    qwb.close()
    records: dict[str, dict[str, Any]] = {}
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        if ws.title in {"Cover Sheet", "Contents", "Notes", "Sections A to S"}:
            continue
        for row_number, row in enumerate(ws.iter_rows(min_row=9, max_col=12, values_only=True), 9):
            sic, year = code(row[0]), number(row[2])
            if not sic or year != 2024:
                continue
            cvs = quality.get(sic, (None, None, None, None))
            records[sic] = financial_row(
                source_id="abs_2024_sections", source_file=str(path.relative_to(ROOT)),
                source_sheet=ws.title, source_row=row_number, industry_code=sic,
                industry_code_level=level(sic), industry_label=str(row[1]).strip(), year=2024,
                geography="United Kingdom", population_basis="ABS reporting units (published workbook labels these as enterprises)",
                reporting_unit_count=number(row[3]), turnover_total_gbp=number(row[4], 1_000_000),
                agva_total_gbp=number(row[5], 1_000_000), purchases_total_gbp=number(row[6], 1_000_000),
                employment_cost_total_gbp=number(row[7], 1_000_000), capital_expenditure_total_gbp=number(row[8], 1_000_000),
                turnover_cv=cvs[0], agva_cv=cvs[1], purchases_cv=cvs[2], employment_cost_cv=cvs[3],
                unit_notes="Monetary workbook values converted from GBP millions to annual GBP.",
                quality_notes="CVs joined by exact SIC and year from ABS national quality measures; null means suppressed/unavailable.",
                provenance_json=provenance("abs_2024_sections", ws.title, row_number, requested_population_label="Number of enterprises", interpreted_population="reporting units", quality_source_id="abs_2024_quality"),
            )
    wb.close()
    return sorted(records.values(), key=lambda row: (len(row["industry_code"]), row["industry_code"]))


def parse_abs_size_bands(manifest: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    path = source_file(manifest, "abs_2024_employment_size_band")
    wb = load_workbook(path, read_only=True, data_only=True)
    records = []
    for ws in wb.worksheets:
        if ws.title not in {"1 to 9", "10 to 49", "50 to 249", "250 and over"}:
            continue
        for row_number, row in enumerate(ws.iter_rows(min_row=9, max_col=13, values_only=True), 9):
            sic, year = code(row[0]), number(row[3])
            if not sic or year != 2024:
                continue
            records.append(financial_row(
                source_id="abs_2024_employment_size_band", source_file=str(path.relative_to(ROOT)), source_sheet=ws.title,
                source_row=row_number, industry_code=sic, industry_code_level=level(sic), industry_label=str(row[1]).strip(),
                year=2024, geography="United Kingdom", population_basis="ABS reporting units by employment size band",
                reporting_unit_count=number(row[4]), turnover_total_gbp=number(row[5], 1_000_000),
                agva_total_gbp=number(row[6], 1_000_000), purchases_total_gbp=number(row[7], 1_000_000),
                employment_cost_total_gbp=number(row[8], 1_000_000), capital_expenditure_total_gbp=number(row[9], 1_000_000),
                employment_size_band=ws.title, unit_notes="Monetary values converted from GBP millions to annual GBP.",
                provenance_json=provenance("abs_2024_employment_size_band", ws.title, row_number),
            ))
    wb.close()
    return records


def parse_uk_business(manifest: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    path = source_file(manifest, "uk_business_2025")
    wb = load_workbook(path, read_only=True, data_only=True)
    records = []
    for sheet, field_name, population in (("Table 2", "enterprise_count", "VAT and/or PAYE based enterprises"), ("Table 17", "local_unit_count", "VAT and/or PAYE based local units")):
        ws = wb[sheet]
        for row_number, row in enumerate(ws.iter_rows(min_row=5, max_col=17, values_only=True), 5):
            sic, label = split_code_label(row[0])
            if not sic or len(sic) != 4:
                continue
            records.append(financial_row(
                source_id="uk_business_2025", source_file=str(path.relative_to(ROOT)), source_sheet=sheet, source_row=row_number,
                industry_code=sic, industry_code_level="class", industry_label=label, year=2025, geography="United Kingdom",
                population_basis=population, **{field_name: number(row[1])},
                unit_notes="Counts are control-rounded to base 5.", provenance_json=provenance("uk_business_2025", sheet, row_number, table_population=population),
            ))
    for sheet, band_type, population in (("Table 4", "employment", "enterprises by employment size band"), ("Table 9 ", "turnover", "enterprises by turnover size band"), ("Table 19", "employment", "local units by employment size band")):
        ws = wb[sheet]
        headers = [str(value).strip() if value is not None else "" for value in next(ws.iter_rows(min_row=5, max_row=5, values_only=True))]
        for row_number, row in enumerate(ws.iter_rows(min_row=6, values_only=True), 6):
            sic, label = split_code_label(row[0])
            if not sic or len(sic) != 4:
                continue
            for column in range(1, min(len(headers), len(row))):
                band = headers[column]
                if not band or band.casefold() == "total":
                    continue
                count_value = number(row[column])
                records.append(financial_row(
                    source_id="uk_business_2025", source_file=str(path.relative_to(ROOT)), source_sheet=sheet, source_row=row_number,
                    industry_code=sic, industry_code_level="class", industry_label=label, year=2025, geography="United Kingdom",
                    population_basis=population, employment_size_band=band if band_type == "employment" else None,
                    turnover_size_band=band if band_type == "turnover" else None, band_count=count_value,
                    unit_notes="Band counts are control-rounded to base 5; turnover bands are GBP thousands.",
                    provenance_json=provenance("uk_business_2025", sheet, row_number, band=band, band_type=band_type, table_population=population),
                ))
    wb.close()
    return records


def parse_bres(manifest: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    path = source_file(manifest, "bres_2024_provisional")
    cv_path = source_file(manifest, "bres_2024_cv")
    labels: dict[str, str] = {}
    wb = load_workbook(path, read_only=True, data_only=True)
    for sheet in ("Division Names", "Group Names", "Sic Names"):
        for row in wb[sheet].iter_rows(values_only=True):
            sic = code(row[0])
            if sic and len(sic) in {2, 3, 5} and row[1]:
                labels[sic] = str(row[1]).strip()
    cv_wb = load_workbook(cv_path, read_only=True, data_only=True)
    cvs: dict[str, float | None] = {}
    for row in cv_wb["Table 2a - GB"].iter_rows(min_row=5, max_col=15, values_only=True):
        sic = next((code(value) for value in row[:3] if code(value)), None)
        if sic:
            cvs[sic] = number(row[14])
    cv_wb.close()
    records = []
    for row_number, row in enumerate(wb["Table 2a GB"].iter_rows(min_row=5, max_col=15, values_only=True), 5):
        sic = next((code(value) for value in row[:3] if code(value)), None)
        if not sic:
            continue
        records.append(employment_row(
            source_id="bres_2024_provisional", source_file=str(path.relative_to(ROOT)), source_sheet="Table 2a GB", source_row=row_number,
            industry_code=sic, industry_code_level=level(sic), industry_label=labels.get(sic), year=2024, geography="Great Britain",
            full_time_total=number(row[5], 1000), part_time_total=number(row[8], 1000), employee_total=number(row[11], 1000),
            employment_total=number(row[14], 1000), employment_cv=cvs.get(sic), population_basis="BRES jobs-based employment estimate",
            unit_notes="Workbook values converted from thousands to persons/jobs; employment includes employees plus working owners.",
            provenance_json=provenance("bres_2024_provisional", "Table 2a GB", row_number, cv_source_id="bres_2024_cv"),
        ))
    wb.close()
    return records


def zip_member(zip_path: Path, contains: str) -> bytes:
    with ZipFile(zip_path) as archive:
        names = [name for name in archive.namelist() if contains in name]
        if len(names) != 1:
            raise RuntimeError(f"Expected one member containing {contains!r} in {zip_path}, found {len(names)}")
        return archive.read(names[0])


def parse_ashe_table5(manifest: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    path = source_file(manifest, "ashe_2025_table5")
    metrics = {"5.5a": "median_hourly_pay_gbp", "5.7a": "median_annual_pay_gbp", "5.9a": "median_weekly_hours"}
    cvs = {"5.5b": "hourly_pay_cv", "5.7b": "annual_pay_cv", "5.9b": "weekly_hours_cv"}
    values: dict[tuple[str, str], dict[str, Any]] = {}
    for token, field in {**metrics, **cvs}.items():
        wb = load_workbook(BytesIO(zip_member(path, token)), read_only=True, data_only=True)
        ws = wb["All"]
        for row_number, row in enumerate(ws.iter_rows(min_row=6, max_col=17, values_only=True), 6):
            raw_code = str(row[1]).strip() if row[1] is not None else ""
            if re.fullmatch(r"[A-Z]\d{8}", raw_code):
                continue
            sic = code(raw_code)
            # This workbook first lists region totals, then UK industry rows.
            # It does not publish a region-by-industry matrix on this sheet.
            if row_number < 22 or not sic or not sic.isdigit() or len(sic) != 2:
                continue
            key = ("United Kingdom", sic)
            item = values.setdefault(key, {"label": str(row[0]).strip(), "row": row_number})
            item[field] = number(row[3])
        wb.close()
    records = []
    for (geography, sic), item in values.items():
        records.append(employment_row(
            source_id="ashe_2025_table5", source_file=str(path.relative_to(ROOT)), source_sheet="All", source_row=item["row"],
            industry_code=sic, industry_code_level="division", industry_label=item["label"], year=2025, geography=geography,
            median_hourly_pay_gbp=item.get("median_hourly_pay_gbp"), median_annual_pay_gbp=item.get("median_annual_pay_gbp"),
            median_weekly_hours=item.get("median_weekly_hours"), hourly_pay_cv=item.get("hourly_pay_cv"),
            annual_pay_cv=item.get("annual_pay_cv"), weekly_hours_cv=item.get("weekly_hours_cv"),
            population_basis="ASHE employee jobs", unit_notes="Medians for all employee jobs; CVs are percentages.",
            provenance_json=provenance("ashe_2025_table5", "All", item["row"], archive_metrics=sorted({**metrics, **cvs})),
        ))
    return records


def parse_ashe_table29(manifest: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    path = source_file(manifest, "ashe_2025_table29")
    metrics = {"29.5a": "median_hourly_pay_gbp", "29.7a": "median_annual_pay_gbp", "29.9a": "median_weekly_hours"}
    cvs = {"29.5b": "hourly_pay_cv", "29.7b": "annual_pay_cv", "29.9b": "weekly_hours_cv"}
    values: dict[tuple[str, str | None], dict[str, Any]] = {}
    for token, field in {**metrics, **cvs}.items():
        wb = load_workbook(BytesIO(zip_member(path, token)), read_only=True, data_only=True)
        ws = wb["All"]
        current_section: str | None = None
        current_label: str | None = None
        for row_number, row in enumerate(ws.iter_rows(min_row=6, max_col=17, values_only=True), 6):
            raw_code = str(row[1]).strip() if row[1] is not None else ""
            if re.fullmatch(r"[A-U]", raw_code, re.I):
                current_section, current_label = raw_code.upper(), str(row[0]).strip()
                key = (current_section, None)
                item = values.setdefault(key, {"label": current_label, "occupation_label": None, "row": row_number})
                item[field] = number(row[3])
                continue
            if not current_section or not re.fullmatch(r"\d{1,4}", raw_code):
                continue
            key = (current_section, raw_code)
            description = str(row[0]).strip()
            occupation_label = description.split(",", 1)[1].strip() if "," in description else description
            item = values.setdefault(key, {"label": current_label, "occupation_label": occupation_label, "row": row_number})
            item[field] = number(row[3])
        wb.close()
    return [employment_row(
        source_id="ashe_2025_table29", source_file=str(path.relative_to(ROOT)), source_sheet="All", source_row=item["row"],
        industry_code=section, industry_code_level="section", industry_label=item["label"], occupation_code=occupation,
        occupation_label=item["occupation_label"], year=2025, geography="United Kingdom",
        median_hourly_pay_gbp=item.get("median_hourly_pay_gbp"), median_annual_pay_gbp=item.get("median_annual_pay_gbp"),
        median_weekly_hours=item.get("median_weekly_hours"), hourly_pay_cv=item.get("hourly_pay_cv"),
        annual_pay_cv=item.get("annual_pay_cv"), weekly_hours_cv=item.get("weekly_hours_cv"),
        population_basis="ASHE employee jobs by industry section and SOC occupation", unit_notes="Medians for all employee jobs; CVs are percentages.",
        provenance_json=provenance("ashe_2025_table29", "All", item["row"], occupation_code=occupation),
    ) for (section, occupation), item in values.items()]


def parse_business_demography(manifest: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    path = source_file(manifest, "business_demography_2024")
    wb = load_workbook(path, read_only=True, data_only=True)
    values: dict[str, dict[str, Any]] = defaultdict(dict)
    for sheet, field in (("Table 1.2", "birth_count"), ("Table 2.2", "death_count"), ("Table 3.2", "active_business_stock")):
        ws = wb[sheet]
        header = list(next(ws.iter_rows(min_row=4, max_row=4, values_only=True)))
        year_column = next(i for i, value in enumerate(header) if str(value).strip() == "2024")
        for row_number, row in enumerate(ws.iter_rows(min_row=5, values_only=True), 5):
            sic, label = split_code_label(row[0])
            if sic and sic.isdigit() and len(sic) in {2, 3}:
                values[sic].update({"label": label, field: number(row[year_column]), "row": row_number})
    ws = wb["Table 5.2a"]
    for row_number, row in enumerate(ws.iter_rows(min_row=5, max_col=12, values_only=True), 5):
        sic, label = split_code_label(row[0])
        if sic and sic.isdigit() and len(sic) in {2, 3}:
            values[sic].update({"label": label, "survival_1y": number(row[3], 0.01), "survival_2y": number(row[5], 0.01),
                                "survival_3y": number(row[7], 0.01), "survival_4y": number(row[9], 0.01),
                                "survival_5y": number(row[11], 0.01), "survival_cohort_year": 2019, "survival_row": row_number})
    wb.close()
    records = []
    for sic, item in values.items():
        active = item.get("active_business_stock")
        births, deaths = item.get("birth_count"), item.get("death_count")
        records.append(risk_row(
            source_id="business_demography_2024", source_file=str(path.relative_to(ROOT)), source_sheet="Tables 1.2/2.2/3.2/5.2a",
            source_row=item.get("row"), industry_code=sic, industry_code_level=level(sic), industry_label=item.get("label"), year=2024,
            geography="United Kingdom", birth_count=births, death_count=deaths, active_business_stock=active,
            birth_rate=(births / active) if births is not None and active else None,
            death_rate=(deaths / active) if deaths is not None and active else None,
            survival_1y=item.get("survival_1y"), survival_2y=item.get("survival_2y"), survival_3y=item.get("survival_3y"),
            survival_4y=item.get("survival_4y"), survival_5y=item.get("survival_5y"), survival_cohort_year=item.get("survival_cohort_year"),
            population_basis="ONS active enterprise population; survival cohort born in 2019",
            unit_notes="Rates stored as fractions from 0 to 1; published survival percentages divided by 100.",
            provenance_json=provenance("business_demography_2024", "Tables 1.2/2.2/3.2/5.2a", item.get("row"), survival_source_row=item.get("survival_row")),
        ))
    return records


def parse_insolvency_demography(manifest: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    path = source_file(manifest, "business_insolvency_demography_2015_2025")
    wb = load_workbook(path, read_only=True, data_only=True)
    counts: dict[str, tuple[str, float | None, int]] = {}
    for row_number, row in enumerate(wb["Table_B1a"].iter_rows(min_row=8, values_only=True), 8):
        sic = code(row[0])
        if sic and sic.isalpha():
            counts[sic] = (str(row[1]).strip(), number(row[52]), row_number)
    records = []
    for row_number, row in enumerate(wb["Table_B1b"].iter_rows(min_row=8, values_only=True), 8):
        sic = code(row[0])
        if not sic or not sic.isalpha():
            continue
        label, count_value, count_row = counts.get(sic, (str(row[1]).strip(), None, row_number))
        records.append(risk_row(
            source_id="business_insolvency_demography_2015_2025", source_file=str(path.relative_to(ROOT)), source_sheet="Table_B1a/Table_B1b",
            source_row=row_number, industry_code=sic, industry_code_level="section", industry_label=label, year=2025,
            geography="England and Wales", insolvency_count=count_value, insolvency_rate_per_10000=number(row[52]),
            insolvency_type="all business insolvencies", population_basis="Insolvency Service business population denominator",
            unit_notes="Published rate per 10,000 businesses; not converted to a fraction in the observation table.",
            provenance_json=provenance("business_insolvency_demography_2015_2025", "Table_B1a/Table_B1b", row_number, count_source_row=count_row, year_column=53),
        ))
    wb.close()
    return records


def parse_current_insolvency_industry(manifest: dict[str, dict[str, str]]) -> list[dict[str, Any]]:
    path = source_file(manifest, "company_insolvencies_industry_2026_06")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Table_A1b"]
    records = []
    for row_number, row in enumerate(ws.iter_rows(min_row=10, values_only=True), 10):
        group, division, section = code(row[2]), code(row[1]), code(row[0])
        sic = group or division or section
        if not sic:
            continue
        records.append(risk_row(
            source_id="company_insolvencies_industry_2026_06", source_file=str(path.relative_to(ROOT)), source_sheet="Table_A1b", source_row=row_number,
            industry_code=sic, industry_code_level=level(sic), industry_label=str(row[3]).strip(), year=2025,
            geography="England and Wales", insolvency_count=number(row[13]), insolvency_type="company insolvencies",
            population_basis="Registered company insolvencies; annual count, not a rate",
            unit_notes="2025 is the latest complete calendar year; 2026 data in the workbook are year-to-date and excluded from priors.",
            provenance_json=provenance("company_insolvencies_industry_2026_06", "Table_A1b", row_number, selected_column="2025"),
        ))
    wb.close()
    return records


def atomic_parquet(records: list[dict[str, Any]], schema: pa.Schema, path: Path, sort_columns: Iterable[str]) -> None:
    if not records:
        raise RuntimeError(f"Refusing to write empty output {path}")
    records.sort(key=lambda row: tuple("" if row.get(column) is None else str(row.get(column)) for column in sort_columns))
    table = pa.Table.from_pylist(records, schema=schema)
    if len(set(table["observation_id"].to_pylist())) != table.num_rows:
        raise RuntimeError(f"Duplicate observation identifiers in {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.name + ".part")
    pq.write_table(table, temp, compression="zstd", use_dictionary=True)
    check = pq.ParquetFile(temp)
    if check.metadata is None or check.metadata.num_rows != len(records):
        raise RuntimeError(f"Parquet validation failed for {path}")
    os.replace(temp, path)


def outputs_valid() -> bool:
    try:
        return all(path.exists() and pq.ParquetFile(path).metadata.num_rows > 0 for path in OUTPUTS.values())
    except Exception:
        return False


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--skip-existing", action="store_true")
    group.add_argument("--force", action="store_true")
    args = parser.parse_args(argv)
    if args.skip_existing and outputs_valid() and not args.force:
        print("All observation outputs already exist and are readable; skipped.")
        return 0
    manifest = load_manifest()
    financial = parse_abs(manifest) + parse_abs_size_bands(manifest) + parse_uk_business(manifest)
    employment = parse_bres(manifest) + parse_ashe_table5(manifest) + parse_ashe_table29(manifest)
    risk = parse_business_demography(manifest) + parse_insolvency_demography(manifest) + parse_current_insolvency_industry(manifest)
    atomic_parquet(financial, FINANCIAL_SCHEMA, OUTPUTS["financial"], ("source_id", "industry_code", "year", "geography", "source_sheet", "source_row"))
    atomic_parquet(employment, EMPLOYMENT_SCHEMA, OUTPUTS["employment"], ("source_id", "industry_code", "occupation_code", "geography"))
    atomic_parquet(risk, RISK_SCHEMA, OUTPUTS["risk"], ("source_id", "industry_code", "year", "geography"))
    print(f"Financial observations: {len(financial):,}")
    print(f"Employment observations: {len(employment):,}")
    print(f"Risk observations: {len(risk):,}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
