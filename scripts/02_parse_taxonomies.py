#!/usr/bin/env python3
"""Parse validated raw taxonomies into deterministic intermediate Parquet tables."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shlex
import sys
import time
import unicodedata
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping

try:
    import pyarrow as pa
    import pyarrow.parquet as pq
    from openpyxl import load_workbook
except ModuleNotFoundError as exc:  # pragma: no cover - startup guidance
    raise SystemExit(
        f"Missing dependency {exc.name!r}. Install requirements.txt before running."
    ) from exc


REPO_ROOT = Path(__file__).resolve().parents[1]
MANIFEST_PATH = REPO_ROOT / "data/source_manifest.csv"
INTERIM_DIR = REPO_ROOT / "data/interim"
REPORT_PATH = REPO_ROOT / "reports/taxonomy_parse_report.md"
PARSER_VERSION = "3"
CHUNK_SIZE = 1024 * 1024

REQUIRED_SOURCE_IDS = (
    "uk_sic_2007_summary_structure",
    "uk_sic_2007_activity_index_2024",
    "voa_scat_revaluation_2026",
    "naics_2022_structure",
    "naics_2022_descriptions",
    "osm_greater_london_pbf",
)

COMMON_FIELDS = [
    pa.field("source", pa.string()),
    pa.field("source_record_id", pa.string()),
    pa.field("source_code", pa.string()),
    pa.field("source_label", pa.string()),
    pa.field("normalised_label", pa.string()),
    pa.field("parent_code", pa.string()),
    pa.field("hierarchy_level", pa.string()),
    pa.field("hierarchy_path", pa.string()),
    pa.field("description", pa.string()),
    pa.field("observed_count", pa.int64()),
    pa.field("physical_venue_candidate", pa.bool_()),
    pa.field("source_file", pa.string()),
    pa.field("source_sheet", pa.string()),
    pa.field("source_row", pa.int64()),
    pa.field("provenance_json", pa.string()),
]
COMMON_FIELD_NAMES = [field.name for field in COMMON_FIELDS]

SIC_LEVEL_ORDER = {
    "Section": 0,
    "Division": 1,
    "Group": 2,
    "Class": 3,
    "Subclass": 4,
}
SIC_LEVEL_WIDTH = {
    "Section": None,
    "Division": 2,
    "Group": 3,
    "Class": 4,
    "Subclass": 5,
}
NAICS_LEVEL_BY_LENGTH = {
    2: "sector",
    3: "subsector",
    4: "industry_group",
    5: "naics_industry",
    6: "national_industry",
}

OSM_KEYS = (
    "amenity",
    "shop",
    "leisure",
    "tourism",
    "office",
    "craft",
    "healthcare",
    "sport",
    "man_made",
    "industrial",
    "landuse",
)
OSM_EXCLUDED_VALUES = {
    "amenity": {
        "bench",
        "bicycle_parking",
        "fire_hydrant",
        "grit_bin",
        "parking_entrance",
        "parking_space",
        "post_box",
        "shelter",
        "waste_basket",
        "waste_disposal",
    },
    "shop": {"closed", "disused", "no", "vacant"},
    "leisure": {"dog_park", "nature_reserve", "park", "picnic_table", "playground"},
    "tourism": {"artwork", "information", "picnic_site", "viewpoint"},
    "office": {"closed", "disused", "no", "vacant"},
    "craft": {"closed", "disused", "no", "vacant"},
    "healthcare": {"no"},
    "sport": {"no"},
    "industrial": {"no"},
}
OSM_ALLOWED_VALUES = {
    "man_made": {"gasometer", "kiln", "silo", "storage_tank", "wastewater_plant", "works"},
    "landuse": {"commercial", "industrial", "retail"},
}
VOA_NON_VENUE_KEYWORDS = (
    "advertising right",
    "communication station",
    "electricity transmission",
    "pipeline",
    "sewer",
    "telecommunication mast",
    "telephone kiosk",
    "water supply network",
)


@dataclass(frozen=True)
class ManifestInput:
    source_id: str
    path: Path
    relative_path: str
    row: dict[str, str]


@dataclass
class ParseBundle:
    records: list[dict[str, Any]]
    schema: pa.Schema
    stats: dict[str, Any]
    sort_key: Callable[[dict[str, Any]], Any]


@dataclass
class ParserResult:
    key: str
    source: str
    input_file: str
    output_path: str
    status: str
    records: int
    duplicates: int
    malformed: int
    unmatched: int
    inspected: list[str]
    counts_by_level: dict[str, int]
    warnings: list[str]
    runtime_seconds: float
    output_size: int
    validation_status: str
    error: str = ""
    extra_stats: dict[str, Any] | None = None


def stable_json(value: Any) -> str:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    )


def stable_id(prefix: str, *parts: Any) -> str:
    payload = stable_json([prefix, *parts]).encode("utf-8")
    return f"{prefix}_{hashlib.sha256(payload).hexdigest()[:24]}"


def normalise_label(value: str | None) -> str | None:
    if value is None:
        return None
    normalised = unicodedata.normalize("NFKC", value)
    return " ".join(normalised.split()).casefold()


def exact_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value)
    return text if text != "" else None


def code_text(value: Any, width: int | None = None) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        text = str(value)
    elif isinstance(value, int):
        text = str(value)
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()
    if not text or text.casefold() in {"na", "n/a", "none"}:
        return None
    if width and text.isdigit():
        return text.zfill(width)
    return text


def numeric_or_none(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text.casefold() in {"[c]", "[z]", "c", "z", "na", "n/a"}:
        return None
    try:
        return float(text.replace(",", ""))
    except ValueError:
        return None


def integer_or_none(value: Any) -> int | None:
    number = numeric_or_none(value)
    return int(number) if number is not None and number.is_integer() else None


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(CHUNK_SIZE):
            digest.update(chunk)
    return digest.hexdigest()


def md5_file(path: Path) -> str:
    digest = hashlib.md5(usedforsecurity=False)
    with path.open("rb") as handle:
        while chunk := handle.read(CHUNK_SIZE):
            digest.update(chunk)
    return digest.hexdigest()


def path_from_manifest(value: str) -> Path:
    candidate = (REPO_ROOT / value).resolve()
    if candidate != REPO_ROOT and REPO_ROOT not in candidate.parents:
        raise ValueError(f"manifest path escapes repository: {value}")
    return candidate


def validate_manifest_inputs() -> tuple[dict[str, ManifestInput], str]:
    if not MANIFEST_PATH.is_file():
        raise FileNotFoundError(f"missing source manifest: {MANIFEST_PATH}")
    with MANIFEST_PATH.open(newline="", encoding="utf-8") as handle:
        rows = {row["source_id"]: row for row in csv.DictReader(handle)}

    inputs: dict[str, ManifestInput] = {}
    for source_id in REQUIRED_SOURCE_IDS:
        if source_id not in rows:
            raise ValueError(f"manifest is missing required source_id {source_id}")
        row = rows[source_id]
        if row.get("access_status") not in {"downloaded", "already_valid"}:
            raise ValueError(f"manifest source {source_id} is not marked accessible")
        if row.get("validation_status") != "passed":
            raise ValueError(f"manifest source {source_id} did not pass acquisition validation")
        relative = row.get("local_path", "")
        path = path_from_manifest(relative)
        if not path.is_file():
            raise FileNotFoundError(f"manifest-listed input is missing: {relative}")
        expected_size = int(row.get("file_size_bytes") or -1)
        if path.stat().st_size != expected_size:
            raise ValueError(
                f"size mismatch for {relative}: manifest={expected_size}, local={path.stat().st_size}"
            )
        expected_sha = row.get("sha256", "").lower()
        actual_sha = sha256_file(path)
        if not expected_sha or actual_sha != expected_sha:
            raise ValueError(f"SHA-256 mismatch for {relative}")
        if source_id == "osm_greater_london_pbf":
            expected_md5 = (row.get("local_md5") or row.get("remote_md5") or "").lower()
            if not expected_md5 or md5_file(path) != expected_md5:
                raise ValueError(f"MD5 mismatch for {relative}")
        inputs[source_id] = ManifestInput(source_id, path, relative, row)

    foursquare = rows.get("foursquare_categories", {}).get("access_status", "not_listed")
    return inputs, foursquare


def common_schema(extra_fields: Iterable[pa.Field]) -> pa.Schema:
    return pa.schema([*COMMON_FIELDS, *extra_fields])


def common_record(
    *,
    source: str,
    source_record_id: str,
    source_code: str | None,
    source_label: str | None,
    parent_code: str | None,
    hierarchy_level: str | None,
    hierarchy_path: str | None,
    description: str | None,
    observed_count: int | None,
    physical_venue_candidate: bool | None,
    source_file: str,
    source_sheet: str | None,
    source_row: int | None,
    provenance: Mapping[str, Any],
) -> dict[str, Any]:
    return {
        "source": source,
        "source_record_id": source_record_id,
        "source_code": source_code,
        "source_label": source_label,
        "normalised_label": normalise_label(source_label),
        "parent_code": parent_code,
        "hierarchy_level": hierarchy_level,
        "hierarchy_path": hierarchy_path,
        "description": description,
        "observed_count": observed_count,
        "physical_venue_candidate": physical_venue_candidate,
        "source_file": source_file,
        "source_sheet": source_sheet,
        "source_row": source_row,
        "provenance_json": stable_json(provenance),
    }


def sic_level(raw_level: Any) -> str | None:
    value = normalise_label(exact_text(raw_level))
    return {
        "section": "Section",
        "division": "Division",
        "group": "Group",
        "class": "Class",
        "sub class": "Subclass",
        "subclass": "Subclass",
    }.get(value or "")


def sic_display_code(level: str, code: str) -> str:
    if level in {"Group", "Class", "Subclass"} and code.isdigit() and len(code) >= 3:
        return f"{code[:2]}.{code[2:]}"
    return code


def sic_code_is_valid(level: str, code: str | None) -> bool:
    if code is None:
        return False
    if level == "Section":
        return bool(re.fullmatch(r"[A-U]", code))
    width = SIC_LEVEL_WIDTH[level]
    return bool(width and re.fullmatch(rf"\d{{{width}}}", code))


def parse_sic_hierarchy(input_file: ManifestInput) -> ParseBundle:
    workbook = load_workbook(input_file.path, read_only=True, data_only=True)
    inspected = list(workbook.sheetnames)
    selected = "reworked structure"
    if selected not in workbook.sheetnames:
        workbook.close()
        raise ValueError("could not locate the official 'reworked structure' worksheet")
    worksheet = workbook[selected]
    records: list[dict[str, Any]] = []
    seen_exact: set[tuple[str, str | None, str | None]] = set()
    duplicates = 0

    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), 2):
        level = sic_level(row[7] if len(row) > 7 else None)
        label = exact_text(row[0] if row else None)
        if not level or label is None:
            continue
        section = code_text(row[1], 1)
        division = code_text(row[2], 2)
        group = code_text(row[3], 3)
        class_code = code_text(row[4], 4)
        subclass = code_text(row[5], 5)
        leaf_code = code_text(row[6], 5)
        codes = {
            "Section": section,
            "Division": division,
            "Group": group,
            "Class": class_code,
            "Subclass": subclass,
        }
        source_code = codes[level]
        exact_key = (level, source_code, label)
        if exact_key in seen_exact:
            duplicates += 1
        seen_exact.add(exact_key)
        parent_code = {
            "Section": None,
            "Division": section,
            "Group": division,
            "Class": group,
            "Subclass": class_code,
        }[level]
        components = []
        for component_level in ("Section", "Division", "Group", "Class", "Subclass"):
            component_code = codes[component_level]
            if component_code:
                components.append(
                    f"{component_level} {sic_display_code(component_level, component_code)}"
                )
            if component_level == level:
                break
        malformed = not sic_code_is_valid(level, source_code)
        record = common_record(
            source="uk_sic_2007",
            source_record_id=stable_id("sic_hierarchy", level, source_code, label),
            source_code=source_code,
            source_label=label,
            parent_code=parent_code,
            hierarchy_level=level,
            hierarchy_path=" > ".join(components),
            description=None,
            observed_count=None,
            physical_venue_candidate=None,
            source_file=input_file.relative_path,
            source_sheet=selected,
            source_row=row_number,
            provenance={
                "original_columns": {
                    "Description": row[0],
                    "SECTION": row[1],
                    "Division": row[2],
                    "Group": row[3],
                    "Class": row[4],
                    "Sub Class": row[5],
                    "Most disaggregated level": row[6],
                    "Level headings": row[7],
                }
            },
        )
        record.update(
            {
                "leaf_code": leaf_code,
                "code_malformed": malformed,
                "orphaned": False,
            }
        )
        records.append(record)
    workbook.close()

    available_codes = {record["source_code"] for record in records if record["source_code"]}
    for record in records:
        record["orphaned"] = bool(
            record["parent_code"] and record["parent_code"] not in available_codes
        )
    stats = {
        "duplicates": duplicates,
        "malformed": sum(record["code_malformed"] for record in records),
        "unmatched": sum(record["orphaned"] for record in records),
        "inspected": inspected,
        "selected_structures": [selected],
        "counts_by_level": dict(sorted(Counter(r["hierarchy_level"] for r in records).items())),
        "warnings": [],
    }
    schema = common_schema(
        [
            pa.field("leaf_code", pa.string()),
            pa.field("code_malformed", pa.bool_()),
            pa.field("orphaned", pa.bool_()),
        ]
    )
    return ParseBundle(records, schema, stats, lambda r: (r["source_row"], r["source_record_id"]))


def load_sic_context(path: Path) -> dict[str, dict[str, Any]]:
    table = pq.read_table(
        path,
        columns=["source_code", "leaf_code", "hierarchy_path", "hierarchy_level"],
    )
    context: dict[str, dict[str, Any]] = {}
    for row in table.to_pylist():
        leaf = row.get("leaf_code")
        if leaf:
            context[leaf] = row
    return context


def parse_sic_activity_index(
    input_file: ManifestInput,
    sic_context: Mapping[str, Mapping[str, Any]],
) -> ParseBundle:
    workbook = load_workbook(input_file.path, read_only=True, data_only=True)
    inspected = list(workbook.sheetnames)
    selected = ("Alphabetical Index", "Numerical Index")
    records: list[dict[str, Any]] = []
    seen_exact: set[tuple[str, str | None, str | None, str]] = set()
    duplicates = 0
    malformed = 0
    unmatched = 0

    for sheet_name in selected:
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ValueError(f"missing required activity index worksheet {sheet_name!r}")
        worksheet = workbook[sheet_name]
        index_type = "alphabetical" if sheet_name.startswith("Alphabetical") else "numerical"
        current_numerical_sic_code: str | None = None
        for row_number, row in enumerate(worksheet.iter_rows(min_row=4, values_only=True), 4):
            activity = exact_text(row[2] if len(row) > 2 else None)
            if activity is None:
                continue
            if index_type == "alphabetical":
                sic_code = code_text(row[0], 5)
                legacy_code = code_text(row[1], 5)
            else:
                first = code_text(row[0], 5)
                second = code_text(row[1], 5)
                if first and re.fullmatch(r"\d{5}", first):
                    current_numerical_sic_code = first
                sic_code = current_numerical_sic_code
                legacy_code = second
            exact_key = (index_type, sic_code, legacy_code, activity)
            if exact_key in seen_exact:
                duplicates += 1
                continue
            seen_exact.add(exact_key)
            is_malformed = not bool(sic_code and re.fullmatch(r"\d{5}", sic_code))
            matched = bool(not is_malformed and sic_code in sic_context)
            malformed += int(is_malformed)
            unmatched += int(not is_malformed and not matched)
            hierarchy = sic_context.get(sic_code or "", {})
            record = common_record(
                source="uk_sic_2007_activity_index",
                source_record_id=stable_id(
                    "sic_activity", index_type, sic_code, legacy_code, activity
                ),
                source_code=sic_code,
                source_label=activity,
                parent_code=hierarchy.get("source_code"),
                hierarchy_level="activity_term",
                hierarchy_path=hierarchy.get("hierarchy_path"),
                description=None,
                observed_count=None,
                physical_venue_candidate=None,
                source_file=input_file.relative_path,
                source_sheet=sheet_name,
                source_row=row_number,
                provenance={
                    "index_type": index_type,
                    "original_columns": {
                        "UK SIC 2007": row[0] if len(row) > 0 else None,
                        "UK SIC 2003": row[1] if len(row) > 1 else None,
                        "Activity": row[2] if len(row) > 2 else None,
                    },
                },
            )
            record.update(
                {
                    "activity_term": activity,
                    "sic_code": sic_code,
                    "legacy_sic_2003_code": legacy_code,
                    "index_type": index_type,
                    "hierarchy_code_matched": matched,
                    "code_malformed": is_malformed,
                }
            )
            records.append(record)
    workbook.close()
    stats = {
        "duplicates": duplicates,
        "malformed": malformed,
        "unmatched": unmatched,
        "inspected": inspected,
        "selected_structures": list(selected),
        "counts_by_level": dict(sorted(Counter(r["index_type"] for r in records).items())),
        "warnings": [],
    }
    schema = common_schema(
        [
            pa.field("activity_term", pa.string()),
            pa.field("sic_code", pa.string()),
            pa.field("legacy_sic_2003_code", pa.string()),
            pa.field("index_type", pa.string()),
            pa.field("hierarchy_code_matched", pa.bool_()),
            pa.field("code_malformed", pa.bool_()),
        ]
    )
    return ParseBundle(
        records,
        schema,
        stats,
        lambda r: (r["index_type"], r["sic_code"] or "", r["activity_term"], r["source_row"]),
    )


def voa_physical_candidate(label: str) -> tuple[bool, str]:
    normalised = normalise_label(label) or ""
    for keyword in VOA_NON_VENUE_KEYWORDS:
        if keyword in normalised:
            return False, f"excluded_keyword:{keyword}"
    return True, "voa_non_domestic_property_default"


def remove_group_suffix(label: str, suffix: str) -> str:
    return re.sub(rf"\s+{re.escape(suffix)}$", "", label, flags=re.IGNORECASE).strip()


def parse_voa_scat(input_file: ManifestInput) -> ParseBundle:
    workbook = load_workbook(input_file.path, read_only=True, data_only=True)
    inspected = list(workbook.sheetnames)
    selected = {
        "RVL_4_1": "England and Wales",
        "RVL_4_2": "England",
        "RVL_4_3": "Wales",
    }
    records: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    duplicates = 0
    malformed = 0
    unmatched = 0

    metric_names = [
        "rateable_properties",
        "rateable_value_2023_gbp_thousands",
        "mean_rateable_value_2023_gbp_thousands",
        "median_rateable_value_2023_gbp_thousands",
        "rateable_value_2026_gbp_thousands",
        "mean_rateable_value_2026_gbp_thousands",
        "median_rateable_value_2026_gbp_thousands",
        "rateable_value_change_gbp_thousands",
        "rateable_value_percentage_change",
    ]
    for sheet_name, geography in selected.items():
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ValueError(f"missing required SCat worksheet {sheet_name}")
        worksheet = workbook[sheet_name]
        header = [exact_text(value) for value in next(worksheet.iter_rows(min_row=6, max_row=6, values_only=True))]
        current_sector: str | None = None
        current_subsector: str | None = None
        for row_number, row in enumerate(worksheet.iter_rows(min_row=7, values_only=True), 7):
            row_type = (exact_text(row[0] if row else None) or "").strip()
            label = exact_text(row[2] if len(row) > 2 else None)
            if row_type == "Sector" and label:
                current_sector = remove_group_suffix(label, "Sector")
                current_subsector = None
                continue
            if row_type == "Sub-sector" and label:
                current_subsector = remove_group_suffix(label, "Sub-sector")
                continue
            if row_type != "SCat" or label is None:
                continue
            scat_code = code_text(row[1], 3)
            exact_key = (sheet_name, scat_code or "")
            if exact_key in seen:
                duplicates += 1
                continue
            seen.add(exact_key)
            is_malformed = not bool(scat_code and re.fullmatch(r"\d{3}", scat_code))
            is_orphaned = current_sector is None
            malformed += int(is_malformed)
            unmatched += int(is_orphaned)
            physical, physical_rule = voa_physical_candidate(label)
            raw_metrics = list(row[3:12])
            suppression = {
                metric_names[index]: str(value)
                for index, value in enumerate(raw_metrics)
                if isinstance(value, str) and value.strip().casefold() in {"[c]", "[z]"}
            }
            values = [numeric_or_none(value) for value in raw_metrics]
            count = integer_or_none(raw_metrics[0])
            path_parts = [part for part in (current_sector, current_subsector) if part]
            path_parts.append(f"SCat {scat_code or '[missing]'}")
            original_columns = {
                column: row[index]
                for index, column in enumerate(header)
                if column is not None and index < len(row)
            }
            record = common_record(
                source="voa_scat_2026",
                source_record_id=stable_id("voa_scat", geography, scat_code, label),
                source_code=scat_code,
                source_label=label,
                parent_code=(
                    f"subsector:{normalise_label(current_subsector)}"
                    if current_subsector
                    else f"sector:{normalise_label(current_sector)}" if current_sector else None
                ),
                hierarchy_level="special_category",
                hierarchy_path=" > ".join(path_parts),
                description=None,
                observed_count=count,
                physical_venue_candidate=physical,
                source_file=input_file.relative_path,
                source_sheet=sheet_name,
                source_row=row_number,
                provenance={
                    "geography": geography,
                    "original_columns": original_columns,
                    "suppression_markers": suppression,
                    "physical_rule": physical_rule,
                },
            )
            record.update(
                {
                    "scat_code": scat_code,
                    "sector": current_sector,
                    "sub_sector": current_subsector,
                    "geography": geography,
                    "rateable_properties": count,
                    "rateable_value_2023_gbp_thousands": values[1],
                    "mean_rateable_value_2023_gbp_thousands": values[2],
                    "median_rateable_value_2023_gbp_thousands": values[3],
                    "rateable_value_2026_gbp_thousands": values[4],
                    "mean_rateable_value_2026_gbp_thousands": values[5],
                    "median_rateable_value_2026_gbp_thousands": values[6],
                    "rateable_value_change_gbp_thousands": values[7],
                    "rateable_value_percentage_change": values[8],
                    "floorspace": None,
                    "physical_rule": physical_rule,
                    "code_malformed": is_malformed,
                    "orphaned": is_orphaned,
                }
            )
            records.append(record)
    workbook.close()
    stats = {
        "duplicates": duplicates,
        "malformed": malformed,
        "unmatched": unmatched,
        "inspected": inspected,
        "selected_structures": list(selected),
        "counts_by_level": dict(sorted(Counter(r["geography"] for r in records).items())),
        "warnings": [
            "The SCat tables provide counts and rateable values but no SCat-level floorspace field; floorspace remains null.",
            "Suppressed [c] and not-applicable [z] metrics remain null and are preserved in provenance_json.",
        ],
        "distinct_scat_codes": len({record["scat_code"] for record in records}),
    }
    schema = common_schema(
        [
            pa.field("scat_code", pa.string()),
            pa.field("sector", pa.string()),
            pa.field("sub_sector", pa.string()),
            pa.field("geography", pa.string()),
            pa.field("rateable_properties", pa.int64()),
            pa.field("rateable_value_2023_gbp_thousands", pa.float64()),
            pa.field("mean_rateable_value_2023_gbp_thousands", pa.float64()),
            pa.field("median_rateable_value_2023_gbp_thousands", pa.float64()),
            pa.field("rateable_value_2026_gbp_thousands", pa.float64()),
            pa.field("mean_rateable_value_2026_gbp_thousands", pa.float64()),
            pa.field("median_rateable_value_2026_gbp_thousands", pa.float64()),
            pa.field("rateable_value_change_gbp_thousands", pa.float64()),
            pa.field("rateable_value_percentage_change", pa.float64()),
            pa.field("floorspace", pa.float64()),
            pa.field("physical_rule", pa.string()),
            pa.field("code_malformed", pa.bool_()),
            pa.field("orphaned", pa.bool_()),
        ]
    )
    return ParseBundle(
        records,
        schema,
        stats,
        lambda r: (r["scat_code"] or "", r["geography"], r["source_row"]),
    )


def naics_level(code: str | None) -> str | None:
    if not code:
        return None
    if re.fullmatch(r"\d{2}-\d{2}", code):
        return "sector"
    if code.isdigit():
        return NAICS_LEVEL_BY_LENGTH.get(len(code))
    return None


def derive_naics_parent(code: str, available_codes: set[str] | None = None) -> str | None:
    level = naics_level(code)
    if level == "sector" or level is None:
        return None
    if len(code) == 3:
        sector = code[:2]
        sector_number = int(sector)
        for range_code, lower, upper in (("31-33", 31, 33), ("44-45", 44, 45), ("48-49", 48, 49)):
            if lower <= sector_number <= upper:
                return range_code
        return sector
    candidate = code[:-1]
    if available_codes is None or candidate in available_codes:
        return candidate
    for length in range(len(code) - 2, 1, -1):
        candidate = code[:length]
        if candidate in available_codes:
            return candidate
    return code[:-1]


def build_naics_path(code: str, parents: Mapping[str, str | None]) -> str:
    parts: list[str] = []
    current: str | None = code
    visited: set[str] = set()
    while current and current not in visited:
        visited.add(current)
        level = naics_level(current) or "unknown"
        parts.append(f"{level.replace('_', ' ').title()} {current}")
        current = parents.get(current)
    return " > ".join(reversed(parts))


def parse_naics_hierarchy(input_file: ManifestInput) -> ParseBundle:
    workbook = load_workbook(input_file.path, read_only=True, data_only=True)
    inspected = list(workbook.sheetnames)
    selected = "2022 NAICS Structure"
    if selected not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"missing required worksheet {selected!r}")
    worksheet = workbook[selected]
    raw_records: list[dict[str, Any]] = []
    seen: set[tuple[str | None, str | None, str | None]] = set()
    duplicates = 0
    for row_number, row in enumerate(worksheet.iter_rows(min_row=4, values_only=True), 4):
        code = code_text(row[1] if len(row) > 1 else None)
        title = exact_text(row[2] if len(row) > 2 else None)
        indicator = exact_text(row[0] if row else None)
        if code is None and title is None:
            continue
        exact_key = (code, title, indicator)
        if exact_key in seen:
            duplicates += 1
            continue
        seen.add(exact_key)
        raw_records.append(
            {
                "code": code,
                "title": title,
                "indicator": indicator,
                "row_number": row_number,
                "raw_row": list(row[:3]),
            }
        )
    workbook.close()
    available_codes = {record["code"] for record in raw_records if record["code"]}
    parents = {
        code: derive_naics_parent(code, available_codes)
        for code in available_codes
        if code is not None
    }
    records: list[dict[str, Any]] = []
    for raw in raw_records:
        code = raw["code"]
        title = raw["title"]
        level = naics_level(code)
        parent = parents.get(code) if code else None
        malformed = level is None or title is None
        orphaned = bool(parent and parent not in available_codes)
        record = common_record(
            source="naics_2022",
            source_record_id=stable_id("naics_hierarchy", code, title),
            source_code=code,
            source_label=title,
            parent_code=parent,
            hierarchy_level=level,
            hierarchy_path=build_naics_path(code, parents) if code else None,
            description=None,
            observed_count=None,
            physical_venue_candidate=None,
            source_file=input_file.relative_path,
            source_sheet=selected,
            source_row=raw["row_number"],
            provenance={
                "original_columns": {
                    "Change Indicator": raw["raw_row"][0],
                    "2022 NAICS Code": raw["raw_row"][1],
                    "2022 NAICS Title": raw["raw_row"][2],
                }
            },
        )
        record.update(
            {
                "change_indicator": raw["indicator"],
                "code_malformed": malformed,
                "orphaned": orphaned,
            }
        )
        records.append(record)
    stats = {
        "duplicates": duplicates,
        "malformed": sum(record["code_malformed"] for record in records),
        "unmatched": sum(record["orphaned"] for record in records),
        "inspected": inspected,
        "selected_structures": [selected],
        "counts_by_level": dict(
            sorted(Counter(r["hierarchy_level"] or "malformed" for r in records).items())
        ),
        "warnings": [],
    }
    schema = common_schema(
        [
            pa.field("change_indicator", pa.string()),
            pa.field("code_malformed", pa.bool_()),
            pa.field("orphaned", pa.bool_()),
        ]
    )
    return ParseBundle(records, schema, stats, lambda r: (r["source_row"], r["source_record_id"]))


def load_naics_context(path: Path) -> dict[str, dict[str, Any]]:
    table = pq.read_table(
        path,
        columns=["source_code", "parent_code", "hierarchy_level", "hierarchy_path"],
    )
    return {row["source_code"]: row for row in table.to_pylist() if row["source_code"]}


def extract_naics_sections(description: str | None) -> tuple[str | None, str | None]:
    if not description:
        return None, None
    lower = description.casefold()
    inclusions: str | None = None
    exclusions: list[str] = []
    examples_marker = "illustrative examples:"
    cross_marker = "cross-references."
    examples_start = lower.find(examples_marker)
    if examples_start >= 0:
        examples_end = lower.find(cross_marker, examples_start)
        if examples_end < 0:
            examples_end = len(description)
        candidate = description[examples_start + len(examples_marker) : examples_end].strip()
        inclusions = candidate or None
    cross_start = lower.find(cross_marker)
    if cross_start >= 0:
        candidate = description[cross_start + len(cross_marker) :].strip()
        if candidate:
            exclusions.append(candidate)
    for match in re.finditer(r"(?i)excluded from[^\n]*(?:\n(?!\n)[^\n]*)*", description):
        candidate = match.group(0).strip()
        if candidate and candidate not in exclusions:
            exclusions.append(candidate)
    return inclusions, "\n\n".join(exclusions) if exclusions else None


def parse_naics_descriptions(
    input_file: ManifestInput,
    naics_context: Mapping[str, Mapping[str, Any]],
) -> ParseBundle:
    workbook = load_workbook(input_file.path, read_only=True, data_only=True)
    inspected = list(workbook.sheetnames)
    selected = "2022_NAICS_Descriptions"
    if selected not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"missing required worksheet {selected!r}")
    worksheet = workbook[selected]
    records: list[dict[str, Any]] = []
    seen: set[tuple[str | None, str | None, str | None]] = set()
    duplicates = 0
    malformed = 0
    unmatched = 0
    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), 2):
        code = code_text(row[0] if row else None)
        title = exact_text(row[1] if len(row) > 1 else None)
        description = exact_text(row[2] if len(row) > 2 else None)
        if code is None and title is None and description is None:
            continue
        exact_key = (code, title, description)
        if exact_key in seen:
            duplicates += 1
            continue
        seen.add(exact_key)
        is_malformed = naics_level(code) is None
        matched = bool(code and code in naics_context)
        malformed += int(is_malformed)
        unmatched += int(not is_malformed and not matched)
        context = naics_context.get(code or "", {})
        inclusions, exclusions = extract_naics_sections(description)
        record = common_record(
            source="naics_2022_descriptions",
            source_record_id=stable_id("naics_description", code, title),
            source_code=code,
            source_label=title,
            parent_code=context.get("parent_code"),
            hierarchy_level=context.get("hierarchy_level") or naics_level(code),
            hierarchy_path=context.get("hierarchy_path"),
            description=description,
            observed_count=None,
            physical_venue_candidate=None,
            source_file=input_file.relative_path,
            source_sheet=selected,
            source_row=row_number,
            provenance={
                "original_columns": {
                    "Code": row[0] if len(row) > 0 else None,
                    "Title": row[1] if len(row) > 1 else None,
                    "Description": row[2] if len(row) > 2 else None,
                }
            },
        )
        record.update(
            {
                "inclusions": inclusions,
                "exclusions": exclusions,
                "hierarchy_code_matched": matched,
                "code_malformed": is_malformed,
            }
        )
        records.append(record)
    workbook.close()
    stats = {
        "duplicates": duplicates,
        "malformed": malformed,
        "unmatched": unmatched,
        "inspected": inspected,
        "selected_structures": [selected],
        "counts_by_level": dict(
            sorted(Counter(r["hierarchy_level"] or "unmatched" for r in records).items())
        ),
        "warnings": [],
    }
    schema = common_schema(
        [
            pa.field("inclusions", pa.string()),
            pa.field("exclusions", pa.string()),
            pa.field("hierarchy_code_matched", pa.bool_()),
            pa.field("code_malformed", pa.bool_()),
        ]
    )
    return ParseBundle(records, schema, stats, lambda r: (r["source_row"], r["source_record_id"]))


def osm_tag_is_relevant(key: str, value: str) -> bool:
    normalised = value.strip().casefold()
    if not normalised:
        return False
    if key in OSM_ALLOWED_VALUES:
        return normalised in OSM_ALLOWED_VALUES[key]
    return normalised not in OSM_EXCLUDED_VALUES.get(key, set())


class OSMTagAggregator:
    """Aggregate selected OSM tags without retaining individual objects."""

    def __init__(self, max_examples: int = 5) -> None:
        self.max_examples = max_examples
        self.combinations: dict[tuple[str, str, str], dict[str, Any]] = {}
        self.objects_inspected = 0
        self.objects_matched = 0

    def add(self, tags: Mapping[str, str], geometry_type: str) -> bool:
        self.objects_inspected += 1
        inspected_tags = {
            key: str(tags[key]).strip()
            for key in OSM_KEYS
            if key in tags and str(tags[key]).strip()
        }
        primary_key: str | None = None
        primary_value: str | None = None
        for key in OSM_KEYS:
            value = inspected_tags.get(key)
            if value and osm_tag_is_relevant(key, value):
                primary_key = key
                primary_value = value
                break
        if primary_key is None or primary_value is None:
            return False
        secondary = {
            key: value
            for key, value in sorted(inspected_tags.items())
            if key != primary_key
        }
        secondary_json = stable_json(secondary)
        key = (primary_key, primary_value, secondary_json)
        aggregate = self.combinations.setdefault(
            key,
            {"count": 0, "names": set(), "geometry_types": set()},
        )
        aggregate["count"] += 1
        aggregate["geometry_types"].add(geometry_type)
        name = exact_text(tags.get("name"))
        if name:
            aggregate["names"].add(name)
            if len(aggregate["names"]) > self.max_examples:
                aggregate["names"] = set(sorted(aggregate["names"])[: self.max_examples])
        self.objects_matched += 1
        return True

    def to_records(self, source_file: str) -> list[dict[str, Any]]:
        records: list[dict[str, Any]] = []
        for (primary_key, primary_value, secondary_json), aggregate in sorted(
            self.combinations.items()
        ):
            label = f"{primary_key}={primary_value}"
            record = common_record(
                source="openstreetmap_greater_london",
                source_record_id=stable_id(
                    "osm_category", primary_key, primary_value, secondary_json
                ),
                source_code=label,
                source_label=label,
                parent_code=primary_key,
                hierarchy_level="tag_combination",
                hierarchy_path=f"OSM > {primary_key} > {primary_value}",
                description=None,
                observed_count=aggregate["count"],
                physical_venue_candidate=True,
                source_file=source_file,
                source_sheet=None,
                source_row=None,
                provenance={
                    "filter_version": PARSER_VERSION,
                    "inspected_tag_keys": list(OSM_KEYS),
                    "secondary_tags": json.loads(secondary_json),
                },
            )
            record["normalised_label"] = normalise_label(
                f"{primary_key} {primary_value.replace('_', ' ')}"
            )
            record.update(
                {
                    "primary_key": primary_key,
                    "primary_value": primary_value,
                    "secondary_tags_json": secondary_json,
                    "example_names": sorted(aggregate["names"]),
                    "geometry_types": sorted(aggregate["geometry_types"]),
                }
            )
            records.append(record)
        return records


def parse_osm_categories(input_file: ManifestInput) -> ParseBundle:
    try:
        import osmium
    except ModuleNotFoundError as exc:
        raise RuntimeError(
            "The Python osmium binding is required for streaming PBF parsing."
        ) from exc

    aggregator = OSMTagAggregator()

    class Handler(osmium.SimpleHandler):
        def _add(self, obj: Any, geometry_type: str) -> None:
            tags = {tag.k: tag.v for tag in obj.tags}
            aggregator.add(tags, geometry_type)

        def node(self, node: Any) -> None:
            self._add(node, "node")

        def way(self, way: Any) -> None:
            self._add(way, "way")

        def relation(self, relation: Any) -> None:
            self._add(relation, "relation")

    Handler().apply_file(input_file.path, locations=False)
    records = aggregator.to_records(input_file.relative_path)
    stats = {
        "duplicates": 0,
        "malformed": 0,
        "unmatched": 0,
        "inspected": [
            "OSM PBF nodes, ways, and relations",
            "tag keys: " + ", ".join(OSM_KEYS),
        ],
        "selected_structures": list(OSM_KEYS),
        "counts_by_level": dict(sorted(Counter(r["primary_key"] for r in records).items())),
        "warnings": [
            "OSM inclusion is tag-rule based; generic and multi-purpose combinations require canonicalisation review."
        ],
        "objects_inspected": aggregator.objects_inspected,
        "physical_venue_observations": aggregator.objects_matched,
        "distinct_tag_combinations": len(records),
    }
    schema = common_schema(
        [
            pa.field("primary_key", pa.string()),
            pa.field("primary_value", pa.string()),
            pa.field("secondary_tags_json", pa.string()),
            pa.field("example_names", pa.list_(pa.string())),
            pa.field("geometry_types", pa.list_(pa.string())),
        ]
    )
    return ParseBundle(
        records,
        schema,
        stats,
        lambda r: (r["primary_key"], r["primary_value"], r["secondary_tags_json"]),
    )


def schema_with_metadata(
    schema: pa.Schema,
    *,
    stats: Mapping[str, Any],
    source_file: str,
) -> pa.Schema:
    metadata = dict(schema.metadata or {})
    metadata.update(
        {
            b"pipeline": b"scripts/02_parse_taxonomies.py",
            b"parser_version": PARSER_VERSION.encode("utf-8"),
            b"source_file": source_file.encode("utf-8"),
            b"parse_stats_json": stable_json(stats).encode("utf-8"),
        }
    )
    return schema.with_metadata(metadata)


def write_parquet_atomic(
    records: list[dict[str, Any]],
    schema: pa.Schema,
    output_path: Path,
    stats: Mapping[str, Any],
    source_file: str,
) -> None:
    if not records:
        raise ValueError(f"refusing to write empty output {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    part_path = output_path.with_name(output_path.name + ".part")
    part_path.unlink(missing_ok=True)
    table = pa.Table.from_pylist(
        records,
        schema=schema_with_metadata(schema, stats=stats, source_file=source_file),
    )
    try:
        pq.write_table(
            table,
            part_path,
            compression="zstd",
            compression_level=9,
            use_dictionary=True,
            write_statistics=True,
            version="2.6",
        )
        validate_parquet_output(part_path)
        os.replace(part_path, output_path)
    except Exception:
        part_path.unlink(missing_ok=True)
        raise


def validate_parquet_output(path: Path) -> tuple[int, dict[str, Any]]:
    parquet_file = pq.ParquetFile(path)
    metadata = parquet_file.metadata
    if metadata is None or metadata.num_rows <= 0:
        raise ValueError(f"Parquet output has no rows: {path}")
    names = set(parquet_file.schema_arrow.names)
    missing = set(COMMON_FIELD_NAMES) - names
    if missing:
        raise ValueError(f"Parquet output lacks common fields {sorted(missing)}: {path}")
    schema_metadata = parquet_file.schema_arrow.metadata or {}
    if schema_metadata.get(b"parser_version", b"").decode() != PARSER_VERSION:
        raise ValueError(f"Parquet output was created by a different parser version: {path}")
    table = pq.read_table(path, columns=["source_record_id"])
    identifiers = table.column("source_record_id").to_pylist()
    if any(identifier is None for identifier in identifiers):
        raise ValueError(f"Parquet output contains null record identifiers: {path}")
    if len(identifiers) != len(set(identifiers)):
        raise ValueError(f"Parquet output contains duplicate record identifiers: {path}")
    try:
        stats = json.loads(schema_metadata[b"parse_stats_json"].decode("utf-8"))
    except (KeyError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError(f"Parquet output lacks readable parse metadata: {path}") from exc
    return metadata.num_rows, stats


def process_parser(
    *,
    key: str,
    source: str,
    input_file: ManifestInput,
    output_path: Path,
    builder: Callable[[], ParseBundle],
    skip_existing: bool,
    force: bool,
) -> ParserResult:
    started = time.perf_counter()
    if output_path.exists() and not force:
        try:
            record_count, stats = validate_parquet_output(output_path)
        except Exception as exc:
            if skip_existing:
                return ParserResult(
                    key,
                    source,
                    input_file.relative_path,
                    output_path.relative_to(REPO_ROOT).as_posix(),
                    "failed",
                    0,
                    0,
                    0,
                    0,
                    [],
                    {},
                    [],
                    time.perf_counter() - started,
                    output_path.stat().st_size,
                    "failed",
                    f"existing output invalid and --skip-existing was supplied: {exc}",
                )
        else:
            return ParserResult(
                key,
                source,
                input_file.relative_path,
                output_path.relative_to(REPO_ROOT).as_posix(),
                "already_valid",
                record_count,
                int(stats.get("duplicates", 0)),
                int(stats.get("malformed", 0)),
                int(stats.get("unmatched", 0)),
                list(stats.get("inspected", [])),
                dict(stats.get("counts_by_level", {})),
                list(stats.get("warnings", [])),
                time.perf_counter() - started,
                output_path.stat().st_size,
                "passed",
                extra_stats=stats,
            )

    try:
        bundle = builder()
        bundle.records.sort(key=bundle.sort_key)
        write_parquet_atomic(
            bundle.records,
            bundle.schema,
            output_path,
            bundle.stats,
            input_file.relative_path,
        )
        record_count, stats = validate_parquet_output(output_path)
        return ParserResult(
            key,
            source,
            input_file.relative_path,
            output_path.relative_to(REPO_ROOT).as_posix(),
            "created",
            record_count,
            int(stats.get("duplicates", 0)),
            int(stats.get("malformed", 0)),
            int(stats.get("unmatched", 0)),
            list(stats.get("inspected", [])),
            dict(stats.get("counts_by_level", {})),
            list(stats.get("warnings", [])),
            time.perf_counter() - started,
            output_path.stat().st_size,
            "passed",
            extra_stats=stats,
        )
    except Exception as exc:
        return ParserResult(
            key,
            source,
            input_file.relative_path,
            output_path.relative_to(REPO_ROOT).as_posix(),
            "failed",
            0,
            0,
            0,
            0,
            [],
            {},
            [],
            time.perf_counter() - started,
            output_path.stat().st_size if output_path.exists() else 0,
            "failed",
            str(exc),
        )


def human_size(size: int) -> str:
    value = float(size)
    for unit in ("B", "KiB", "MiB", "GiB"):
        if value < 1024 or unit == "GiB":
            return f"{value:.0f} {unit}" if unit == "B" else f"{value:.2f} {unit}"
        value /= 1024
    return f"{size} B"


def markdown(value: Any) -> str:
    return str(value).replace("|", "\\|").replace("\n", "<br>")


def write_parse_report(
    results: list[ParserResult],
    *,
    foursquare_status: str,
    command: str,
    inputs_validated: int,
) -> None:
    by_key = {result.key: result for result in results}
    total_records = sum(result.records for result in results if result.validation_status == "passed")
    warnings = [warning for result in results for warning in result.warnings]
    failures = [result for result in results if result.validation_status != "passed"]
    lines = [
        "# Taxonomy parse report",
        "",
        f"Parser version: `{PARSER_VERSION}`",
        f"Exact command: `{command}`",
        f"Manifest-validated required inputs: {inputs_validated}",
        f"Foursquare status: `{foursquare_status}` (optional; not parsed)",
        "",
        "| Source | Input file | Input sheets or structures inspected | Records parsed | Duplicates | Malformed records | Unmatched codes | Output path | Output size | Runtime | Validation status |",
        "|---|---|---|---:|---:|---:|---:|---|---:|---:|---|",
    ]
    for result in results:
        lines.append(
            "| "
            + " | ".join(
                markdown(value)
                for value in (
                    result.source,
                    f"`{result.input_file}`",
                    ", ".join(result.inspected) or "—",
                    result.records,
                    result.duplicates,
                    result.malformed,
                    result.unmatched,
                    f"`{result.output_path}`",
                    human_size(result.output_size) if result.output_size else "—",
                    f"{result.runtime_seconds:.3f}s",
                    result.validation_status if not result.error else f"failed: {result.error}",
                )
            )
            + " |"
        )

    def count(key: str) -> int:
        return by_key.get(key).records if key in by_key else 0

    osm_stats = (by_key.get("osm").extra_stats or {}) if by_key.get("osm") else {}
    voa_stats = (by_key.get("voa_scat").extra_stats or {}) if by_key.get("voa_scat") else {}
    lines.extend(
        [
            "",
            "## Record totals",
            "",
            f"- UK SIC hierarchy records: {count('sic_hierarchy'):,}",
            f"- UK SIC activity terms: {count('sic_activity'):,}",
            f"- VOA SCat records: {count('voa_scat'):,} across {int(voa_stats.get('distinct_scat_codes', 0)):,} distinct SCat codes",
            f"- NAICS hierarchy records: {count('naics_hierarchy'):,}",
            f"- NAICS description records: {count('naics_descriptions'):,}",
            f"- Distinct OSM tag combinations: {count('osm'):,}",
            f"- OSM physical venue observations: {int(osm_stats.get('physical_venue_observations', 0)):,}",
            f"- Total intermediate records: {total_records:,}",
            "",
            "## Records by hierarchy level or partition",
            "",
        ]
    )
    for result in results:
        detail = ", ".join(
            f"{level}={value:,}" for level, value in sorted(result.counts_by_level.items())
        )
        lines.append(f"- {result.source}: {detail or 'not available'}")

    lines.extend(["", "## Warnings", ""])
    if warnings:
        lines.extend(f"- {warning}" for warning in dict.fromkeys(warnings))
    else:
        lines.append("None.")
    lines.extend(["", "## Parser failures", ""])
    if failures:
        lines.extend(f"- {result.source}: {result.error}" for result in failures)
    else:
        lines.append("None.")
    lines.extend(
        [
            "",
            "## Ambiguities for canonicalisation",
            "",
            "- UK SIC activity codes are five-digit leaf codes; unsplit hierarchy classes retain their official four-digit class code plus a separate five-digit `leaf_code`.",
            "- Alphabetical and numerical SIC index records remain separate through `index_type`; only exact within-index duplicates are removed.",
            "- VOA emits one record per SCat and geography (England and Wales, England, Wales). Suppressed metrics remain null rather than being estimated.",
            "- NAICS title text, including the source workbook's trilateral `T` marker, is preserved exactly.",
            "- OSM combinations reflect observed tagging practice, not equivalence; generic, secondary, and multi-purpose tags require review before merging.",
            "- Foursquare remains optional and was not awaited.",
            "",
        ]
    )
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    part_path = REPORT_PATH.with_name(REPORT_PATH.name + ".part")
    part_path.write_text("\n".join(lines), encoding="utf-8")
    os.replace(part_path, REPORT_PATH)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    group = parser.add_mutually_exclusive_group()
    group.add_argument(
        "--skip-existing",
        action="store_true",
        help="revalidate and retain valid parser-version-matched outputs",
    )
    group.add_argument(
        "--force",
        action="store_true",
        help="rebuild and atomically replace every intermediate output",
    )
    return parser.parse_args()


def main() -> int:
    if sys.version_info < (3, 11):
        raise SystemExit("Python 3.11 or newer is required")
    args = parse_args()
    INTERIM_DIR.mkdir(parents=True, exist_ok=True)
    command = shlex.join([sys.executable, *sys.argv])
    try:
        inputs, foursquare_status = validate_manifest_inputs()
    except Exception as exc:
        print(f"Input validation failed: {exc}", file=sys.stderr)
        return 1
    print(f"Manifest validation passed for {len(inputs)} required inputs")

    results: list[ParserResult] = []
    sic_output = INTERIM_DIR / "sic_hierarchy.parquet"
    result = process_parser(
        key="sic_hierarchy",
        source="UK SIC hierarchy",
        input_file=inputs["uk_sic_2007_summary_structure"],
        output_path=sic_output,
        builder=lambda: parse_sic_hierarchy(inputs["uk_sic_2007_summary_structure"]),
        skip_existing=args.skip_existing,
        force=args.force,
    )
    results.append(result)
    print(f"UK SIC hierarchy: {result.status} ({result.records:,} records)")

    sic_context: dict[str, dict[str, Any]] = {}
    if result.validation_status == "passed":
        sic_context = load_sic_context(sic_output)
    result = process_parser(
        key="sic_activity",
        source="UK SIC activity index",
        input_file=inputs["uk_sic_2007_activity_index_2024"],
        output_path=INTERIM_DIR / "sic_activity_index.parquet",
        builder=lambda: parse_sic_activity_index(
            inputs["uk_sic_2007_activity_index_2024"], sic_context
        ),
        skip_existing=args.skip_existing,
        force=args.force,
    )
    results.append(result)
    print(f"UK SIC activity index: {result.status} ({result.records:,} records)")

    result = process_parser(
        key="voa_scat",
        source="VOA Special Categories",
        input_file=inputs["voa_scat_revaluation_2026"],
        output_path=INTERIM_DIR / "voa_scat.parquet",
        builder=lambda: parse_voa_scat(inputs["voa_scat_revaluation_2026"]),
        skip_existing=args.skip_existing,
        force=args.force,
    )
    results.append(result)
    print(f"VOA Special Categories: {result.status} ({result.records:,} records)")

    naics_output = INTERIM_DIR / "naics_hierarchy.parquet"
    result = process_parser(
        key="naics_hierarchy",
        source="NAICS hierarchy",
        input_file=inputs["naics_2022_structure"],
        output_path=naics_output,
        builder=lambda: parse_naics_hierarchy(inputs["naics_2022_structure"]),
        skip_existing=args.skip_existing,
        force=args.force,
    )
    results.append(result)
    print(f"NAICS hierarchy: {result.status} ({result.records:,} records)")

    naics_context: dict[str, dict[str, Any]] = {}
    if result.validation_status == "passed":
        naics_context = load_naics_context(naics_output)
    result = process_parser(
        key="naics_descriptions",
        source="NAICS descriptions",
        input_file=inputs["naics_2022_descriptions"],
        output_path=INTERIM_DIR / "naics_descriptions.parquet",
        builder=lambda: parse_naics_descriptions(
            inputs["naics_2022_descriptions"], naics_context
        ),
        skip_existing=args.skip_existing,
        force=args.force,
    )
    results.append(result)
    print(f"NAICS descriptions: {result.status} ({result.records:,} records)")

    result = process_parser(
        key="osm",
        source="Greater London OpenStreetMap",
        input_file=inputs["osm_greater_london_pbf"],
        output_path=INTERIM_DIR / "osm_business_categories.parquet",
        builder=lambda: parse_osm_categories(inputs["osm_greater_london_pbf"]),
        skip_existing=args.skip_existing,
        force=args.force,
    )
    results.append(result)
    print(f"Greater London OSM: {result.status} ({result.records:,} records)")

    write_parse_report(
        results,
        foursquare_status=foursquare_status,
        command=command,
        inputs_validated=len(inputs),
    )
    failures = [result for result in results if result.validation_status != "passed"]
    print(f"Report: {REPORT_PATH.relative_to(REPO_ROOT)}")
    if failures:
        print(f"Parser failures: {len(failures)}", file=sys.stderr)
        return 1
    print(f"Validated intermediate outputs: {len(results)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
