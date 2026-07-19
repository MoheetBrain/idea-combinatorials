#!/usr/bin/env python3
"""Download and validate the authoritative phase-one taxonomy source files."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shlex
import shutil
import subprocess
import sys
import time
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Callable
from urllib.parse import urlparse

try:
    import requests
except ModuleNotFoundError as exc:  # pragma: no cover - startup guidance
    raise SystemExit(
        "Missing dependency 'requests'. Install requirements.txt before running."
    ) from exc


REPO_ROOT = Path(__file__).resolve().parents[1]
MANIFEST_PATH = REPO_ROOT / "data/source_manifest.csv"
DOWNLOAD_REPORT_PATH = REPO_ROOT / "reports/source_download_report.md"
MANUAL_REPORT_PATH = REPO_ROOT / "reports/manual_actions_required.md"
FOURSQUARE_DIR = REPO_ROOT / "data/raw/foursquare/categories"

USER_AGENT = (
    "idea-combinatorials-phase-one-source-downloader/1.0 "
    "(+https://github.com/; authoritative taxonomy acquisition)"
)
CONNECT_TIMEOUT_SECONDS = 20
READ_TIMEOUT_SECONDS = 180
MAX_ATTEMPTS = 3
CHUNK_SIZE = 1024 * 1024
TRANSIENT_HTTP_STATUSES = {408, 425, 429, 500, 502, 503, 504}
PBF_MINIMUM_PLAUSIBLE_SIZE = 1_000_000

MANIFEST_FIELDS = [
    "source_id",
    "dataset_name",
    "provider",
    "source_page_url",
    "requested_download_url",
    "resolved_download_url",
    "local_path",
    "file_format",
    "file_size_bytes",
    "sha256",
    "remote_md5",
    "local_md5",
    "retrieved_at_utc",
    "http_status",
    "validation_status",
    "validation_details",
    "licence",
    "access_status",
    "notes",
]


@dataclass(frozen=True)
class SourceSpec:
    source_id: str
    dataset_name: str
    provider: str
    source_page_url: str
    download_url: str
    local_path: str
    file_format: str
    licence: str
    allowed_domains: tuple[str, ...]


PUBLIC_SOURCES = [
    SourceSpec(
        source_id="uk_sic_2007_summary_structure",
        dataset_name="UK SIC 2007 hierarchy",
        provider="Office for National Statistics",
        source_page_url=(
            "https://www.ons.gov.uk/methodology/classificationsandstandards/"
            "ukstandardindustrialclassificationofeconomicactivities/uksic2007"
        ),
        download_url=(
            "https://www.ons.gov.uk/file?uri=%2Fmethodology%2Fclassificationsandstandards%2F"
            "ukstandardindustrialclassificationofeconomicactivities%2Fuksic2007%2F"
            "publisheduksicsummaryofstructureworksheet.xlsx"
        ),
        local_path="data/raw/sic/uk_sic_2007_summary_structure.xlsx",
        file_format="xlsx",
        licence="Open Government Licence v3.0",
        allowed_domains=("ons.gov.uk",),
    ),
    SourceSpec(
        source_id="uk_sic_2007_activity_index_2024",
        dataset_name="UK SIC 2007 detailed index",
        provider="Office for National Statistics",
        source_page_url=(
            "https://www.ons.gov.uk/methodology/classificationsandstandards/"
            "ukstandardindustrialclassificationofeconomicactivities/uksic2007"
        ),
        download_url=(
            "https://www.ons.gov.uk/file?uri=%2Fmethodology%2Fclassificationsandstandards%2F"
            "ukstandardindustrialclassificationofeconomicactivities%2Fuksic2007%2F"
            "uksic2007indexeswithaddendumdecember2022newformatsept20241.xlsx"
        ),
        local_path="data/raw/sic/uk_sic_2007_activity_index_2024.xlsx",
        file_format="xlsx",
        licence="Open Government Licence v3.0",
        allowed_domains=("ons.gov.uk",),
    ),
    SourceSpec(
        source_id="voa_scat_revaluation_2026",
        dataset_name="VOA Special Category physical-property data",
        provider="Valuation Office Agency",
        source_page_url=(
            "https://www.gov.uk/government/statistics/non-domestic-rating-change-in-"
            "rateable-value-of-rating-lists-england-and-wales-2026-revaluation-draft-list/"
            "non-domestic-rating-reval-2026-draft-list-and-background-information"
        ),
        download_url=(
            "https://assets.publishing.service.gov.uk/media/692971dd9c1eda2cdf034269/"
            "ndr_reval_2026.xlsx"
        ),
        local_path="data/raw/voa/voa_scat_revaluation_2026.xlsx",
        file_format="xlsx",
        licence="Open Government Licence v3.0",
        allowed_domains=("publishing.service.gov.uk", "gov.uk"),
    ),
    SourceSpec(
        source_id="naics_2022_structure",
        dataset_name="NAICS 2022 hierarchy",
        provider="U.S. Census Bureau",
        source_page_url="https://www.census.gov/naics/",
        download_url="https://www.census.gov/naics/2022NAICS/2022_NAICS_Structure.xlsx",
        local_path="data/raw/naics/naics_2022_structure.xlsx",
        file_format="xlsx",
        licence="U.S. Census Bureau source",
        allowed_domains=("census.gov",),
    ),
    SourceSpec(
        source_id="naics_2022_descriptions",
        dataset_name="NAICS 2022 detailed descriptions",
        provider="U.S. Census Bureau",
        source_page_url="https://www.census.gov/naics/",
        download_url="https://www.census.gov/naics/2022NAICS/2022_NAICS_Descriptions.xlsx",
        local_path="data/raw/naics/naics_2022_descriptions.xlsx",
        file_format="xlsx",
        licence="U.S. Census Bureau source",
        allowed_domains=("census.gov",),
    ),
    SourceSpec(
        source_id="osm_greater_london_md5",
        dataset_name="Greater London OpenStreetMap extract checksum",
        provider="Geofabrik / OpenStreetMap contributors",
        source_page_url=(
            "https://download.geofabrik.de/europe/united-kingdom/england/greater-london.html"
        ),
        download_url=(
            "https://download.geofabrik.de/europe/united-kingdom/england/"
            "greater-london-latest.osm.pbf.md5"
        ),
        local_path="data/raw/osm/greater-london-latest.osm.pbf.md5",
        file_format="md5",
        licence="Open Database Licence 1.0; attribution required",
        allowed_domains=("geofabrik.de",),
    ),
    SourceSpec(
        source_id="osm_greater_london_pbf",
        dataset_name="Greater London OpenStreetMap extract",
        provider="Geofabrik / OpenStreetMap contributors",
        source_page_url=(
            "https://download.geofabrik.de/europe/united-kingdom/england/greater-london.html"
        ),
        download_url=(
            "https://download.geofabrik.de/europe/united-kingdom/england/"
            "greater-london-latest.osm.pbf"
        ),
        local_path="data/raw/osm/greater-london-latest.osm.pbf",
        file_format="osm.pbf",
        licence="Open Database Licence 1.0; attribution required",
        allowed_domains=("geofabrik.de",),
    ),
]


class DownloadFailure(RuntimeError):
    def __init__(
        self,
        message: str,
        *,
        status: int | None = None,
        resolved_url: str = "",
        transient: bool = False,
    ) -> None:
        super().__init__(message)
        self.status = status
        self.resolved_url = resolved_url
        self.transient = transient


def utc_now() -> str:
    return datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def relative_path(path: Path) -> str:
    try:
        return path.resolve().relative_to(REPO_ROOT).as_posix()
    except ValueError:
        return str(path)


def ensure_directories() -> None:
    for path in (
        REPO_ROOT / "data/raw/sic",
        REPO_ROOT / "data/raw/voa",
        REPO_ROOT / "data/raw/naics",
        REPO_ROOT / "data/raw/osm",
        FOURSQUARE_DIR,
        REPO_ROOT / "scripts",
        REPO_ROOT / "reports",
    ):
        path.mkdir(parents=True, exist_ok=True)


def digest_file(path: Path, algorithm: str) -> str:
    digest = hashlib.new(algorithm)
    with path.open("rb") as handle:
        while chunk := handle.read(CHUNK_SIZE):
            digest.update(chunk)
    return digest.hexdigest()


def payload_prefix(path: Path, length: int = 4096) -> bytes:
    with path.open("rb") as handle:
        return handle.read(length).lstrip().lower()


def reject_error_payload(path: Path) -> None:
    if not path.exists() or path.stat().st_size == 0:
        raise ValueError("file is empty")
    prefix = payload_prefix(path)
    if prefix.startswith((b"<!doctype html", b"<html", b"<head", b"<body")):
        raise ValueError("file contains HTML rather than the requested data format")
    if prefix.startswith((b"{", b"[")):
        try:
            json.loads(prefix.decode("utf-8", errors="strict"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            pass
        else:
            raise ValueError("file contains a JSON response rather than the requested data format")


def validate_xlsx(path: Path) -> dict[str, Any]:
    reject_error_payload(path)
    if not payload_prefix(path, 4).startswith(b"pk") or not zipfile.is_zipfile(path):
        raise ValueError("not a valid ZIP-based Office workbook")
    with zipfile.ZipFile(path) as archive:
        corrupt_member = archive.testzip()
    if corrupt_member:
        raise ValueError(f"workbook ZIP member failed its CRC check: {corrupt_member}")
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise RuntimeError("openpyxl is required to validate Excel files") from exc

    workbook_handle = path.open("rb")
    try:
        # A file object is intentional: temporary downloads end in `.xlsx.part`,
        # while openpyxl otherwise rejects them based only on the final suffix.
        workbook = load_workbook(workbook_handle, read_only=True, data_only=True)
    except Exception as exc:
        workbook_handle.close()
        raise ValueError(f"openpyxl could not parse workbook: {exc}") from exc

    try:
        if not workbook.worksheets:
            raise ValueError("workbook contains no worksheets")
        sheets: list[dict[str, Any]] = []
        workbook_has_content = False
        for worksheet in workbook.worksheets:
            nonempty_cells = 0
            last_row = 0
            last_column = 0
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is not None and str(cell.value).strip() != "":
                        nonempty_cells += 1
                        last_row = max(last_row, cell.row)
                        last_column = max(last_column, cell.column)
            if nonempty_cells:
                workbook_has_content = True
            sheets.append(
                {
                    "name": worksheet.title,
                    "approx_rows": last_row,
                    "approx_columns": last_column,
                    "nonempty_cells": nonempty_cells,
                }
            )
        if not workbook_has_content:
            raise ValueError("no worksheet contains non-empty cells")
        return {"validator": "zipfile+openpyxl", "worksheets": sheets}
    finally:
        workbook.close()
        workbook_handle.close()


def read_remote_md5(path: Path) -> str:
    reject_error_payload(path)
    text = path.read_text(encoding="utf-8", errors="strict")
    match = re.search(r"(?i)(?<![0-9a-f])([0-9a-f]{32})(?![0-9a-f])", text)
    if not match:
        raise ValueError("checksum file does not contain a 32-character MD5 hash")
    return match.group(1).lower()


def validate_md5_file(path: Path) -> dict[str, Any]:
    return {"validator": "MD5 text parser", "expected_md5": read_remote_md5(path)}


def validate_pbf(path: Path, expected_md5: str) -> dict[str, Any]:
    reject_error_payload(path)
    size = path.stat().st_size
    if size < PBF_MINIMUM_PLAUSIBLE_SIZE:
        raise ValueError(
            f"PBF size {size:,} bytes is below the plausibility threshold "
            f"of {PBF_MINIMUM_PLAUSIBLE_SIZE:,} bytes"
        )
    local_md5 = digest_file(path, "md5")
    if not expected_md5:
        raise ValueError("official remote MD5 is unavailable")
    if local_md5.lower() != expected_md5.lower():
        raise ValueError(f"MD5 mismatch: expected {expected_md5}, calculated {local_md5}")

    details: dict[str, Any] = {
        "validator": "official MD5 comparison",
        "expected_md5": expected_md5.lower(),
        "local_md5": local_md5.lower(),
        "plausible_size": True,
    }
    osmium = shutil.which("osmium")
    if osmium:
        completed = subprocess.run(
            [osmium, "fileinfo", str(path)],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
        if completed.returncode != 0:
            raise ValueError(
                "osmium fileinfo failed: "
                + (completed.stderr.strip() or f"exit status {completed.returncode}")
            )
        details["osmium_fileinfo"] = "passed"
        first_line = next((line.strip() for line in completed.stdout.splitlines() if line.strip()), "")
        if first_line:
            details["osmium_summary"] = first_line
    else:
        details["osmium_fileinfo"] = "not performed (osmium not installed; optional)"
    return details


def validate_parquet(path: Path) -> dict[str, Any]:
    reject_error_payload(path)
    try:
        import pyarrow.parquet as parquet
    except ModuleNotFoundError as exc:
        raise RuntimeError("pyarrow is required to validate Parquet files") from exc
    try:
        parquet_file = parquet.ParquetFile(path)
        metadata = parquet_file.metadata
        schema = parquet_file.schema_arrow
    except Exception as exc:
        raise ValueError(f"pyarrow could not read Parquet metadata: {exc}") from exc
    if metadata is None or metadata.num_rows <= 0:
        raise ValueError("Parquet metadata reports no rows")
    if not schema.names:
        raise ValueError("Parquet schema contains no columns")
    return {
        "validator": "pyarrow.parquet metadata",
        "rows": metadata.num_rows,
        "row_groups": metadata.num_row_groups,
        "columns": schema.names,
    }


def validate_file(spec: SourceSpec, path: Path, expected_md5: str = "") -> dict[str, Any]:
    validators: dict[str, Callable[[Path], dict[str, Any]]] = {
        "xlsx": validate_xlsx,
        "md5": validate_md5_file,
        "parquet": validate_parquet,
    }
    if spec.file_format == "osm.pbf":
        return validate_pbf(path, expected_md5)
    try:
        validator = validators[spec.file_format]
    except KeyError as exc:
        raise ValueError(f"no validator configured for {spec.file_format}") from exc
    return validator(path)


def domain_allowed(url: str, allowed_domains: tuple[str, ...]) -> bool:
    parsed = urlparse(url)
    if parsed.scheme != "https" or not parsed.hostname:
        return False
    hostname = parsed.hostname.lower()
    return any(hostname == domain or hostname.endswith("." + domain) for domain in allowed_domains)


def stream_download(
    session: requests.Session,
    spec: SourceSpec,
    part_path: Path,
    *,
    headers: dict[str, str] | None = None,
) -> tuple[str, int]:
    last_error: Exception | None = None
    for attempt in range(1, MAX_ATTEMPTS + 1):
        part_path.unlink(missing_ok=True)
        try:
            with session.get(
                spec.download_url,
                stream=True,
                allow_redirects=True,
                timeout=(CONNECT_TIMEOUT_SECONDS, READ_TIMEOUT_SECONDS),
                headers=headers,
            ) as response:
                resolved_url = response.url
                status = response.status_code
                if status >= 400:
                    raise DownloadFailure(
                        f"HTTP {status} for {spec.download_url}",
                        status=status,
                        resolved_url=resolved_url,
                        transient=status in TRANSIENT_HTTP_STATUSES,
                    )
                if not domain_allowed(resolved_url, spec.allowed_domains):
                    raise DownloadFailure(
                        f"redirect resolved to an unrelated domain: {resolved_url}",
                        status=status,
                        resolved_url=resolved_url,
                    )
                content_type = response.headers.get("Content-Type", "").lower()
                if "text/html" in content_type or "application/json" in content_type:
                    raise DownloadFailure(
                        f"unexpected response content type {content_type!r}",
                        status=status,
                        resolved_url=resolved_url,
                    )

                bytes_written = 0
                with part_path.open("wb") as handle:
                    for chunk in response.iter_content(chunk_size=CHUNK_SIZE):
                        if chunk:
                            handle.write(chunk)
                            bytes_written += len(chunk)
                    handle.flush()
                    os.fsync(handle.fileno())
                if bytes_written == 0:
                    raise DownloadFailure(
                        "server returned an empty file",
                        status=status,
                        resolved_url=resolved_url,
                    )
                content_length = response.headers.get("Content-Length")
                content_encoding = response.headers.get("Content-Encoding", "identity").lower()
                if content_length and content_encoding in ("", "identity"):
                    expected_length = int(content_length)
                    if bytes_written != expected_length:
                        raise DownloadFailure(
                            f"truncated download: expected {expected_length} bytes, "
                            f"received {bytes_written}",
                            status=status,
                            resolved_url=resolved_url,
                            transient=True,
                        )
                reject_error_payload(part_path)
                return resolved_url, status
        except DownloadFailure as exc:
            last_error = exc
            retry = exc.transient
        except (requests.ConnectionError, requests.Timeout, requests.ChunkedEncodingError) as exc:
            last_error = exc
            retry = True
        except requests.RequestException as exc:
            last_error = exc
            retry = False
        except (OSError, ValueError) as exc:
            last_error = exc
            retry = False

        part_path.unlink(missing_ok=True)
        if not retry or attempt == MAX_ATTEMPTS:
            break
        delay = 2 ** (attempt - 1)
        print(f"  transient failure; retrying in {delay}s ({attempt}/{MAX_ATTEMPTS})")
        time.sleep(delay)

    if isinstance(last_error, DownloadFailure):
        raise last_error
    raise DownloadFailure(f"download failed: {last_error}") from last_error


def empty_row(spec: SourceSpec) -> dict[str, str]:
    row = {field: "" for field in MANIFEST_FIELDS}
    row.update(
        {
            "source_id": spec.source_id,
            "dataset_name": spec.dataset_name,
            "provider": spec.provider,
            "source_page_url": spec.source_page_url,
            "requested_download_url": spec.download_url,
            "local_path": spec.local_path,
            "file_format": spec.file_format,
            "licence": spec.licence,
        }
    )
    return row


def load_previous_manifest() -> dict[str, dict[str, str]]:
    if not MANIFEST_PATH.exists():
        return {}
    try:
        with MANIFEST_PATH.open(newline="", encoding="utf-8") as handle:
            return {row["source_id"]: row for row in csv.DictReader(handle) if row.get("source_id")}
    except (OSError, csv.Error, KeyError):
        return {}


def prior_retrieval_time(prior: dict[str, str] | None, path: Path) -> str:
    if prior and prior.get("retrieved_at_utc"):
        return prior["retrieved_at_utc"]
    timestamp = datetime.fromtimestamp(path.stat().st_mtime, UTC).replace(microsecond=0)
    return timestamp.isoformat().replace("+00:00", "Z")


def successful_row(
    spec: SourceSpec,
    path: Path,
    details: dict[str, Any],
    *,
    access_status: str,
    resolved_url: str,
    http_status: str | int,
    retrieved_at: str,
    expected_md5: str = "",
    note: str = "",
) -> dict[str, str]:
    row = empty_row(spec)
    local_md5 = digest_file(path, "md5") if spec.file_format == "osm.pbf" else ""
    row.update(
        {
            "resolved_download_url": resolved_url,
            "file_size_bytes": str(path.stat().st_size),
            "sha256": digest_file(path, "sha256"),
            "remote_md5": expected_md5 if spec.file_format == "osm.pbf" else "",
            "local_md5": local_md5,
            "retrieved_at_utc": retrieved_at,
            "http_status": str(http_status),
            "validation_status": "passed",
            "validation_details": json.dumps(details, ensure_ascii=False, sort_keys=True),
            "access_status": access_status,
            "notes": note,
        }
    )
    return row


def failed_row(
    spec: SourceSpec,
    error: Exception,
    *,
    access_status: str = "failed",
    note: str = "",
) -> dict[str, str]:
    row = empty_row(spec)
    if isinstance(error, DownloadFailure):
        row["http_status"] = str(error.status or "")
        row["resolved_download_url"] = error.resolved_url
    row.update(
        {
            "validation_status": "failed" if access_status == "failed" else "not_performed",
            "validation_details": str(error),
            "access_status": access_status,
            "notes": note,
        }
    )
    return row


def process_public_source(
    session: requests.Session,
    spec: SourceSpec,
    previous: dict[str, dict[str, str]],
    *,
    skip_existing: bool,
    force: bool,
    expected_md5: str = "",
) -> dict[str, str]:
    path = REPO_ROOT / spec.local_path
    part_path = path.with_name(path.name + ".part")
    prior = previous.get(spec.source_id)
    path.parent.mkdir(parents=True, exist_ok=True)

    if path.exists() and not force:
        try:
            details = validate_file(spec, path, expected_md5)
        except Exception as exc:
            if skip_existing:
                return failed_row(
                    spec,
                    exc,
                    note="Existing file is invalid; --skip-existing prevented replacement.",
                )
            print(f"  existing file is invalid and will be replaced after a valid download: {exc}")
        else:
            return successful_row(
                spec,
                path,
                details,
                access_status="already_valid",
                resolved_url=(prior or {}).get("resolved_download_url", spec.download_url),
                http_status=(prior or {}).get("http_status", ""),
                retrieved_at=prior_retrieval_time(prior, path),
                expected_md5=expected_md5,
                note=(
                    "Existing file revalidated; no download performed because --skip-existing was supplied."
                    if skip_existing
                    else "Existing file revalidated; valid files require --force to be replaced."
                ),
            )

    try:
        resolved_url, status = stream_download(session, spec, part_path)
        details = validate_file(spec, part_path, expected_md5)
        os.replace(part_path, path)
        return successful_row(
            spec,
            path,
            details,
            access_status="downloaded",
            resolved_url=resolved_url,
            http_status=status,
            retrieved_at=utc_now(),
            expected_md5=expected_md5,
            note="Downloaded to a .part file and moved atomically after validation.",
        )
    except Exception as exc:
        part_path.unlink(missing_ok=True)
        return failed_row(spec, exc)


def get_hugging_face_token() -> tuple[str | None, str]:
    try:
        from huggingface_hub import get_token
    except ModuleNotFoundError:
        return None, "huggingface_hub is not installed"
    token = get_token()
    if not token:
        portal_hint = any(
            os.environ.get(name)
            for name in ("FSQ_API_TOKEN", "FOURSQUARE_API_TOKEN", "FOURSQUARE_TOKEN")
        )
        if portal_hint:
            return None, (
                "A Foursquare token variable is present, but Hugging Face authentication is not "
                "available; portal connection metadata must be configured manually."
            )
        return None, "No HF_TOKEN or authenticated Hugging Face CLI session was detected."
    return token, "Hugging Face credentials detected without exposing the token."


def redact_provider_error(error: Exception, token: str) -> str:
    """Return provider diagnostics without ever reproducing a credential."""
    message = str(error).replace(token, "[REDACTED]")
    return re.sub(r"(?i)bearer\s+[a-z0-9._~-]+", "Bearer [REDACTED]", message)


def foursquare_manual_row(reason: str) -> dict[str, str]:
    spec = SourceSpec(
        source_id="foursquare_categories",
        dataset_name="Foursquare category taxonomy (latest release)",
        provider="Foursquare / Hugging Face",
        source_page_url="https://docs.foursquare.com/data-products/docs/access-fsq-os-places",
        download_url="https://huggingface.co/datasets/foursquare/fsq-os-places",
        local_path="data/raw/foursquare/categories/",
        file_format="parquet",
        licence="Apache License 2.0, subject to current access terms",
        allowed_domains=("huggingface.co",),
    )
    return failed_row(
        spec,
        RuntimeError(reason),
        access_status="manual_required",
        note="No unofficial mirror or substitute was used.",
    )


def download_hugging_face_file(
    spec: SourceSpec,
    token: str,
    path_in_repo: str,
    destination: Path,
    cache_dir: Path,
) -> None:
    try:
        from huggingface_hub import hf_hub_download
    except ModuleNotFoundError as exc:
        raise RuntimeError("huggingface_hub is required for gated downloads") from exc

    last_error: Exception | None = None
    for attempt in range(1, MAX_ATTEMPTS + 1):
        try:
            cached_path = Path(
                hf_hub_download(
                    repo_id="foursquare/fsq-os-places",
                    filename=path_in_repo,
                    repo_type="dataset",
                    revision="main",
                    token=token,
                    cache_dir=cache_dir,
                )
            )
            part_path = destination.with_name(destination.name + ".part")
            part_path.unlink(missing_ok=True)
            with cached_path.open("rb") as source, part_path.open("wb") as target:
                shutil.copyfileobj(source, target, length=CHUNK_SIZE)
                target.flush()
                os.fsync(target.fileno())
            return
        except Exception as exc:  # hub exceptions differ across supported releases
            last_error = exc
            if attempt == MAX_ATTEMPTS:
                break
            delay = 2 ** (attempt - 1)
            print(f"  Hugging Face download retry in {delay}s ({attempt}/{MAX_ATTEMPTS})")
            time.sleep(delay)
    safe_error = redact_provider_error(last_error, token) if last_error else "unknown provider error"
    raise DownloadFailure(f"Hugging Face download failed after three attempts: {safe_error}")


def process_foursquare(
    previous: dict[str, dict[str, str]],
    *,
    skip_existing: bool,
    force: bool,
) -> tuple[list[dict[str, str]], str]:
    token, auth_note = get_hugging_face_token()
    print(f"Foursquare authentication check: {auth_note}")
    if not token:
        return [foursquare_manual_row(auth_note)], auth_note

    try:
        from huggingface_hub import HfApi, hf_hub_url
        api = HfApi()
        identity = api.whoami(token=token)
        if not identity:
            raise RuntimeError("Hugging Face authentication status could not be confirmed")
        repo_files = api.list_repo_files(
            repo_id="foursquare/fsq-os-places",
            repo_type="dataset",
            token=token,
        )
    except Exception as exc:
        reason = (
            "Hugging Face authentication exists, but the gated dataset could not be listed. "
            "Accept the dataset access terms and retry. Provider response: "
            f"{redact_provider_error(exc, token)}"
        )
        return [foursquare_manual_row(reason)], reason

    pattern = re.compile(
        r"^release/dt=(\d{4}-\d{2}-\d{2})/categories/parquet/([^/]+\.parquet)$",
        re.IGNORECASE,
    )
    matches = [(match.group(1), path, match.group(2)) for path in repo_files if (match := pattern.match(path))]
    if not matches:
        reason = (
            "The authenticated repository listing contained no files matching "
            "release/dt=YYYY-MM-DD/categories/parquet/*.parquet."
        )
        return [foursquare_manual_row(reason)], reason

    latest_release = max(date for date, _, _ in matches)
    selected = sorted((path, filename) for date, path, filename in matches if date == latest_release)
    rows: list[dict[str, str]] = []
    cache_dir = REPO_ROOT / "data/raw/foursquare/.hf_download_cache"
    try:
        for index, (path_in_repo, filename) in enumerate(selected, start=1):
            source_id = "foursquare_categories_" + re.sub(r"[^a-z0-9]+", "_", filename.lower()).strip("_")
            url = hf_hub_url(
                repo_id="foursquare/fsq-os-places",
                filename=path_in_repo,
                repo_type="dataset",
                revision="main",
            )
            spec = SourceSpec(
                source_id=source_id,
                dataset_name=f"Foursquare category taxonomy ({latest_release})",
                provider="Foursquare / Hugging Face",
                source_page_url="https://huggingface.co/datasets/foursquare/fsq-os-places",
                download_url=url,
                local_path=f"data/raw/foursquare/categories/{filename}",
                file_format="parquet",
                licence="Apache License 2.0, subject to current access terms",
                allowed_domains=("huggingface.co", "hf.co", "xethub.hf.co"),
            )
            destination = REPO_ROOT / spec.local_path
            prior = previous.get(source_id)
            print(f"Foursquare categories [{index}/{len(selected)}]: {filename}")
            if destination.exists() and not force:
                same_release = bool(prior and prior.get("requested_download_url") == url)
                try:
                    details = validate_parquet(destination)
                    if not same_release and prior:
                        raise ValueError("existing file belongs to a different recorded release")
                except Exception as exc:
                    if skip_existing:
                        rows.append(
                            failed_row(
                                spec,
                                exc,
                                note="Existing file is invalid or stale; --skip-existing prevented replacement.",
                            )
                        )
                        continue
                else:
                    rows.append(
                        successful_row(
                            spec,
                            destination,
                            details,
                            access_status="already_valid",
                            resolved_url=(prior or {}).get("resolved_download_url", url),
                            http_status=(prior or {}).get("http_status", ""),
                            retrieved_at=prior_retrieval_time(prior, destination),
                            note=f"Latest release {latest_release}; existing Parquet file revalidated.",
                        )
                    )
                    continue

            part_path = destination.with_name(destination.name + ".part")
            try:
                download_hugging_face_file(spec, token, path_in_repo, destination, cache_dir)
                details = validate_parquet(part_path)
                os.replace(part_path, destination)
                rows.append(
                    successful_row(
                        spec,
                        destination,
                        details,
                        access_status="downloaded",
                        resolved_url=url,
                        http_status=200,
                        retrieved_at=utc_now(),
                        note=f"Newest repository release selected: {latest_release}.",
                    )
                )
            except Exception as exc:
                part_path.unlink(missing_ok=True)
                rows.append(failed_row(spec, exc))
    finally:
        shutil.rmtree(cache_dir, ignore_errors=True)

    status = (
        f"Downloaded/validated {sum(row['access_status'] in ('downloaded', 'already_valid') for row in rows)} "
        f"category file(s) from release {latest_release}."
    )
    return rows, status


def write_manifest(rows: list[dict[str, str]]) -> None:
    temp_path = MANIFEST_PATH.with_name(MANIFEST_PATH.name + ".part")
    with temp_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=MANIFEST_FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    os.replace(temp_path, MANIFEST_PATH)


def human_size(size: int) -> str:
    value = float(size)
    for unit in ("B", "KiB", "MiB", "GiB", "TiB"):
        if value < 1024 or unit == "TiB":
            return f"{value:.0f} {unit}" if unit == "B" else f"{value:.2f} {unit}"
        value /= 1024
    return f"{size} B"


def markdown_escape(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", " ")


def report_warning(row: dict[str, str]) -> str:
    warnings: list[str] = []
    if row["access_status"] == "manual_required":
        warnings.append("Manual authentication required")
    if row["access_status"] == "failed":
        warnings.append(row["validation_details"] or "Failed")
    if "not performed (osmium not installed" in row["validation_details"]:
        warnings.append("Optional osmium validation not performed")
    return "; ".join(warnings) or "None"


def write_manual_report(foursquare_status: str, rows: list[dict[str, str]]) -> None:
    manual_rows = [row for row in rows if row["access_status"] == "manual_required"]
    lines = ["# Manual actions required", ""]
    if not manual_rows:
        lines.extend(["None. Foursquare category access and validation completed.", ""])
    else:
        lines.extend(
            [
                "## Foursquare category taxonomy",
                "",
                markdown_escape(foursquare_status),
                "",
                "The dataset is gated. No bypass, mirror, substitute, or placeholder was used.",
                "",
                "1. Sign in at https://huggingface.co/datasets/foursquare/fsq-os-places.",
                "2. Review and accept the repository's current gated-access terms.",
                "3. Create a read-only Hugging Face access token at https://huggingface.co/settings/tokens.",
                "4. Authenticate locally with `hf auth login` (or `huggingface-cli login` on older clients).",
                "5. Rerun `python scripts/01_download_sources.py --skip-existing`.",
                "",
                "Do not paste a token into a report, shell history, source file, or command argument. "
                "`HF_TOKEN` may instead be supplied through a secure environment mechanism.",
                "",
                "Alternatively, create an account at https://places.foursquare.com and follow the "
                "portal's generated Iceberg connection instructions. The repository script uses the "
                "Hugging Face route because it can enumerate and select only the latest category Parquet files.",
                "",
            ]
        )
    MANUAL_REPORT_PATH.write_text("\n".join(lines), encoding="utf-8")


def write_download_report(
    rows: list[dict[str, str]],
    command: str,
    foursquare_status: str,
) -> None:
    successful = [row for row in rows if row["access_status"] in ("downloaded", "already_valid")]
    failed = [row for row in rows if row["access_status"] == "failed"]
    manual = [row for row in rows if row["access_status"] == "manual_required"]
    total_size = sum(int(row["file_size_bytes"] or 0) for row in successful)
    lines = [
        "# Phase-one source download report",
        "",
        f"Generated at: `{utc_now()}`",
        "",
        "| Dataset | Status | Saved path | File size | Validation performed | SHA-256 | Warnings |",
        "|---|---|---|---:|---|---|---|",
    ]
    for row in rows:
        validation = (
            row["validation_details"] if row["validation_status"] != "passed" else row["validation_details"]
        )
        sha = row["sha256"] or "—"
        lines.append(
            "| "
            + " | ".join(
                markdown_escape(value)
                for value in (
                    row["dataset_name"],
                    row["access_status"],
                    f"`{row['local_path']}`",
                    human_size(int(row["file_size_bytes"] or 0)) if row["file_size_bytes"] else "—",
                    validation or "not performed",
                    f"`{sha}`" if sha != "—" else sha,
                    report_warning(row),
                )
            )
            + " |"
        )

    lines.extend(["", "## Successful downloads", ""])
    if successful:
        lines.extend(f"- `{row['local_path']}` — {row['access_status']}" for row in successful)
    else:
        lines.append("None.")
    lines.extend(["", "## Failed downloads", ""])
    if failed:
        lines.extend(
            f"- `{row['local_path']}` — {markdown_escape(row['validation_details'])}" for row in failed
        )
    else:
        lines.append("None.")
    lines.extend(["", "## Manual actions required", ""])
    if manual:
        lines.append(f"- Foursquare: {markdown_escape(foursquare_status)}")
        lines.append("- See `reports/manual_actions_required.md` for exact login steps.")
    else:
        lines.append("None.")
    lines.extend(
        [
            "",
            "## Totals",
            "",
            f"- Successful validated files: {len(successful)}",
            f"- Failed files: {len(failed)}",
            f"- Manual-required sources: {len(manual)}",
            f"- Total disk usage of successful source files: {human_size(total_size)} ({total_size} bytes)",
            "",
            "## Exact command used",
            "",
            f"`{command}`",
            "",
        ]
    )
    DOWNLOAD_REPORT_PATH.write_text("\n".join(lines), encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    group = parser.add_mutually_exclusive_group()
    group.add_argument(
        "--skip-existing",
        action="store_true",
        help="revalidate and retain valid files; do not replace invalid existing files",
    )
    group.add_argument(
        "--force",
        action="store_true",
        help="redownload and atomically replace files, including valid existing files",
    )
    return parser.parse_args()


def main() -> int:
    if sys.version_info < (3, 11):
        raise SystemExit("Python 3.11 or newer is required")
    args = parse_args()
    ensure_directories()
    previous = load_previous_manifest()
    command = shlex.join([sys.executable, *sys.argv])
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT, "Accept-Encoding": "identity"})

    rows: list[dict[str, str]] = []
    remote_md5 = ""
    print("Downloading and validating authoritative public sources")
    for index, spec in enumerate(PUBLIC_SOURCES, start=1):
        print(f"[{index}/{len(PUBLIC_SOURCES)}] {spec.dataset_name}")
        expected_md5 = remote_md5 if spec.file_format == "osm.pbf" else ""
        if spec.file_format == "osm.pbf" and not expected_md5:
            row = failed_row(
                spec,
                RuntimeError("official MD5 checksum was unavailable or failed validation"),
            )
        else:
            row = process_public_source(
                session,
                spec,
                previous,
                skip_existing=args.skip_existing,
                force=args.force,
                expected_md5=expected_md5,
            )
        rows.append(row)
        print(f"  {row['access_status']}: {row['validation_status']}")
        if spec.file_format == "md5" and row["validation_status"] == "passed":
            try:
                remote_md5 = read_remote_md5(REPO_ROOT / spec.local_path)
            except Exception:
                remote_md5 = ""

    foursquare_rows, foursquare_status = process_foursquare(
        previous,
        skip_existing=args.skip_existing,
        force=args.force,
    )
    rows.extend(foursquare_rows)
    write_manifest(rows)
    write_manual_report(foursquare_status, rows)
    write_download_report(rows, command, foursquare_status)

    public_failures = [
        row for row in rows if row["source_id"] in {spec.source_id for spec in PUBLIC_SOURCES} and row["access_status"] == "failed"
    ]
    successful = sum(row["access_status"] in ("downloaded", "already_valid") for row in rows)
    print(f"Validated source files: {successful}")
    print(f"Manifest: {relative_path(MANIFEST_PATH)}")
    print(f"Report: {relative_path(DOWNLOAD_REPORT_PATH)}")
    if public_failures:
        print(f"Required public source failures: {len(public_failures)}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
